#!/usr/bin/env python3
"""Structural QA for the generated VCC PowerPoint."""

from __future__ import annotations

import re
import sys
import zipfile
from pathlib import Path

from pptx import Presentation


DECK = Path(__file__).with_name("VCC_2026_Three_Person_Team_Strategy_EN.pptx")
CJK = re.compile(r"[\u3400-\u4dbf\u4e00-\u9fff\uf900-\ufaff]")


def iter_text(shape):
    if getattr(shape, "has_text_frame", False):
        yield shape.text
    if getattr(shape, "has_table", False):
        for row in shape.table.rows:
            for cell in row.cells:
                yield cell.text
    if getattr(shape, "shape_type", None) == 6:  # group
        for child in shape.shapes:
            yield from iter_text(child)


def main():
    errors = []
    warnings = []
    if not DECK.exists():
        errors.append(f"missing deck: {DECK}")
        print("\n".join(errors))
        return 1

    with zipfile.ZipFile(DECK) as zf:
        bad = zf.testzip()
        if bad:
            errors.append(f"bad zip member: {bad}")

    prs = Presentation(DECK)
    if len(prs.slides) != 46:
        errors.append(f"expected 46 slides, found {len(prs.slides)}")
    ratio = prs.slide_width / prs.slide_height
    if abs(ratio - 16 / 9) > 0.01:
        errors.append(f"unexpected aspect ratio: {ratio:.4f}")

    slide_w = prs.slide_width
    slide_h = prs.slide_height
    all_text = []
    min_font_pt = 999.0
    title_texts = []

    for idx, slide in enumerate(prs.slides, start=1):
        texts = []
        for shape in slide.shapes:
            if shape.left < -1000 or shape.top < -1000:
                errors.append(f"slide {idx}: negative shape origin")
            if shape.left + shape.width > slide_w + 15000 or shape.top + shape.height > slide_h + 15000:
                errors.append(f"slide {idx}: shape outside slide bounds")
            for value in iter_text(shape):
                if value:
                    texts.append(value)
                    all_text.append(value)
            if getattr(shape, "has_text_frame", False):
                for p in shape.text_frame.paragraphs:
                    for run in p.runs:
                        if run.font.size:
                            min_font_pt = min(min_font_pt, run.font.size.pt)
        if not texts:
            errors.append(f"slide {idx}: no text")
        # The first non-empty all-caps item can be a section label; record the longest top-level title candidate.
        candidates = [t for t in texts if 4 <= len(t) <= 100 and "\n" not in t]
        if candidates:
            title_texts.append(max(candidates, key=len))

    joined = "\n".join(all_text)
    if CJK.search(joined):
        errors.append("CJK characters found in slide content; deck must be English-only")
    if min_font_pt < 7.0:
        warnings.append(f"minimum explicit font size is {min_font_pt:.1f} pt")

    # Check notes and slide XML are present for every slide.
    with zipfile.ZipFile(DECK) as zf:
        names = set(zf.namelist())
        slide_xml = [n for n in names if re.fullmatch(r"ppt/slides/slide\d+\.xml", n)]
        notes_xml = [n for n in names if re.fullmatch(r"ppt/notesSlides/notesSlide\d+\.xml", n)]
        if len(slide_xml) != 46:
            errors.append(f"expected 46 slide XML files, found {len(slide_xml)}")
        if len(notes_xml) != 46:
            warnings.append(f"expected 46 notes XML files, found {len(notes_xml)}")

    print(f"deck={DECK}")
    print(f"bytes={DECK.stat().st_size}")
    print(f"slides={len(prs.slides)}")
    print(f"aspect_ratio={ratio:.4f}")
    print(f"minimum_explicit_font_pt={min_font_pt:.1f}")
    print(f"text_characters={len(joined)}")
    print(f"errors={len(errors)}")
    for item in errors:
        print(f"ERROR: {item}")
    print(f"warnings={len(warnings)}")
    for item in warnings:
        print(f"WARNING: {item}")
    return 1 if errors else 0


if __name__ == "__main__":
    sys.exit(main())
