#!/usr/bin/env python3
"""Approximate native-shape PowerPoint renderer for visual QA contact sheets.

This is not a PowerPoint rendering engine. It renders the subset of editable shapes
used by the generated deck closely enough to catch gross layout and overflow issues.
"""

from __future__ import annotations

import math
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN


HERE = Path(__file__).resolve().parent
DECK = HERE / "VCC_2026_Three_Person_Team_Strategy_EN.pptx"
OUTDIR = HERE / "preview"
REGULAR = "/usr/share/fonts/julietaula-montserrat/Montserrat-Regular.otf"
BOLD = "/usr/share/fonts/julietaula-montserrat/Montserrat-Bold.otf"
MONO = "/usr/share/fonts/adobe-source-code-pro/SourceCodePro-Regular.otf"
EMU = 914400
PPI = 120


def px(value):
    return int(round(value / EMU * PPI))


def rgb(value, default=(0, 0, 0)):
    try:
        raw = str(value.rgb)
        if len(raw) == 6:
            return tuple(int(raw[i:i + 2], 16) for i in (0, 2, 4))
    except Exception:
        pass
    return default


def shape_fill(shape, default=None):
    try:
        return rgb(shape.fill.fore_color, default or (0, 0, 0))
    except Exception:
        return default


def shape_line(shape, default=None):
    try:
        return rgb(shape.line.color, default or (95, 110, 140))
    except Exception:
        return default


def font_for(run, pt):
    name = (run.font.name or "") if run is not None else ""
    bold = bool(run.font.bold) if run is not None else False
    path = MONO if "Code" in name else (BOLD if bold else REGULAR)
    return ImageFont.truetype(path, max(7, int(pt * PPI / 72)))


def wrap(draw, value, font, max_w):
    out = []
    for raw_line in value.split("\n"):
        words = raw_line.split(" ")
        if not words:
            out.append("")
            continue
        current = ""
        for word in words:
            trial = word if not current else current + " " + word
            if draw.textlength(trial, font=font) <= max_w or not current:
                current = trial
            else:
                out.append(current)
                current = word
        out.append(current)
    return out


def render_text(draw, shape, x0, y0, x1, y1):
    tf = shape.text_frame
    paragraphs = []
    for p in tf.paragraphs:
        value = p.text
        if not value:
            continue
        run = p.runs[0] if p.runs else None
        pt = (run.font.size.pt if run is not None and run.font.size else 12)
        font = font_for(run, pt)
        color = rgb(run.font.color, (247, 249, 252)) if run is not None else (247, 249, 252)
        align = p.alignment or PP_ALIGN.LEFT
        paragraphs.append((value, font, color, align, max(1, int(pt * PPI / 72 * 1.20))))
    if not paragraphs:
        return
    ml = px(tf.margin_left or 0)
    mr = px(tf.margin_right or 0)
    mt = px(tf.margin_top or 0)
    mb = px(tf.margin_bottom or 0)
    avail_w = max(4, x1 - x0 - ml - mr)
    laid = []
    total_h = 0
    for value, font, color, align, line_h in paragraphs:
        lines = wrap(draw, value, font, avail_w)
        laid.append((lines, font, color, align, line_h))
        total_h += line_h * len(lines) + max(1, int(line_h * 0.08))
    anchor = tf.vertical_anchor
    if anchor == MSO_ANCHOR.MIDDLE:
        yy = y0 + max(mt, (y1 - y0 - total_h) // 2)
    elif anchor == MSO_ANCHOR.BOTTOM:
        yy = y1 - mb - total_h
    else:
        yy = y0 + mt
    for lines, font, color, align, line_h in laid:
        for value in lines:
            width = draw.textlength(value, font=font)
            if align == PP_ALIGN.CENTER:
                xx = x0 + ml + (avail_w - width) / 2
            elif align == PP_ALIGN.RIGHT:
                xx = x1 - mr - width
            else:
                xx = x0 + ml
            draw.text((xx, yy), value, font=font, fill=color)
            yy += line_h
        yy += max(1, int(line_h * 0.08))


def render_shape(draw, shape):
    x0, y0 = px(shape.left), px(shape.top)
    x1, y1 = x0 + px(shape.width), y0 + px(shape.height)
    fill = shape_fill(shape)
    outline = shape_line(shape)
    width = 1
    try:
        width = max(1, int(shape.line.width.pt * PPI / 72)) if shape.line.width else 1
    except Exception:
        pass

    if shape.shape_type == MSO_SHAPE_TYPE.LINE:
        draw.line((x0, y0, x1, y1), fill=outline or (95, 110, 140), width=width)
        return

    if shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE:
        kind = shape.auto_shape_type
        if kind == MSO_SHAPE.OVAL:
            draw.ellipse((x0, y0, x1, y1), fill=fill, outline=outline, width=width)
        elif kind == MSO_SHAPE.ROUNDED_RECTANGLE:
            draw.rounded_rectangle((x0, y0, x1, y1), radius=max(3, int((y1 - y0) * 0.15)),
                                   fill=fill, outline=outline, width=width)
        elif kind == MSO_SHAPE.ISOSCELES_TRIANGLE:
            # Orientation is not material for layout QA.
            draw.polygon(((x0 + x1) / 2, y0, x0, y1, x1, y1), fill=fill, outline=outline)
        else:
            draw.rectangle((x0, y0, x1, y1), fill=fill, outline=outline, width=width)
    elif fill is not None:
        draw.rectangle((x0, y0, x1, y1), fill=fill)

    if getattr(shape, "has_text_frame", False) and shape.text.strip():
        render_text(draw, shape, x0, y0, x1, y1)


def render_slide(prs, slide, index):
    width = int(round(prs.slide_width / EMU * PPI))
    height = int(round(prs.slide_height / EMU * PPI))
    bg = (8, 16, 31)
    image = Image.new("RGB", (width, height), bg)
    draw = ImageDraw.Draw(image)
    for shape in slide.shapes:
        render_shape(draw, shape)
    path = OUTDIR / f"slide_{index:02d}.png"
    image.save(path, optimize=True)
    return image


def contact_sheet(images, start_index, sheet_index):
    thumb_w, thumb_h = 400, 225
    label_h = 24
    canvas = Image.new("RGB", (thumb_w * 4, (thumb_h + label_h) * 3), (18, 27, 46))
    draw = ImageDraw.Draw(canvas)
    font = ImageFont.truetype(BOLD, 13)
    for i, image in enumerate(images):
        row, col = divmod(i, 4)
        x, y = col * thumb_w, row * (thumb_h + label_h)
        canvas.paste(image.resize((thumb_w, thumb_h), Image.Resampling.LANCZOS), (x, y))
        draw.text((x + 8, y + thumb_h + 4), f"Slide {start_index + i:02d}", font=font, fill=(169, 181, 204))
    path = OUTDIR / f"contact_sheet_{sheet_index:02d}.png"
    canvas.save(path, optimize=True)
    return path


def main():
    OUTDIR.mkdir(parents=True, exist_ok=True)
    prs = Presentation(DECK)
    images = [render_slide(prs, slide, i) for i, slide in enumerate(prs.slides, start=1)]
    sheets = []
    for i in range(0, len(images), 12):
        sheets.append(contact_sheet(images[i:i + 12], i + 1, i // 12 + 1))
    print(f"slides_rendered={len(images)}")
    for path in sheets:
        print(f"contact_sheet={path}")


if __name__ == "__main__":
    main()
