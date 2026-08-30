#!/usr/bin/env python3
"""Build the English VCC 2026 discussion deck with editable PowerPoint shapes."""

from __future__ import annotations

import math
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "presentation" / "VCC_2026_Three_Person_Team_Strategy_EN.pptx"
NOTES_OUT = ROOT / "presentation" / "VCC_2026_Speaker_Notes_EN.md"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
prs.core_properties.title = "Virtual Cell Challenge 2026 — Three-Person Team Strategy"
prs.core_properties.subject = "Zero-shot CRISPRi Perturb-seq prediction"
prs.core_properties.author = "Virtual Cell Challenge team"
prs.core_properties.keywords = "VCC 2026, single-cell, CRISPRi, Perturb-seq, diffusion, reinforcement learning"

W, H = 13.333, 7.5

# Theme
BG = "08101F"
BG2 = "0D172A"
SURFACE = "111C31"
SURFACE2 = "172640"
WHITE = "F7F9FC"
MUTED = "A9B5CC"
MUTED2 = "71809B"
PURPLE = "8B5CF6"
PURPLE2 = "C4B5FD"
CYAN = "22D3EE"
CYAN2 = "A5F3FC"
GREEN = "34D399"
GREEN2 = "A7F3D0"
PINK = "F472B6"
PINK2 = "FBCFE8"
YELLOW = "FBBF24"
YELLOW2 = "FDE68A"
RED = "FB7185"
BLUE = "60A5FA"

FONT = "Montserrat"
MONO = "Source Code Pro"

SLIDES = []
NOTES = []


def C(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def set_fill(shape, color: str, transparency: int = 0):
    shape.fill.solid()
    shape.fill.fore_color.rgb = C(color)
    shape.fill.transparency = transparency


def set_line(shape, color: str, width: float = 1.0, transparency: int = 0):
    shape.line.color.rgb = C(color)
    shape.line.width = Pt(width)
    shape.line.transparency = transparency


def rect(slide, x, y, w, h, fill=SURFACE, line=None, radius=True, transparency=0):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    set_fill(shape, fill, transparency)
    if line:
        set_line(shape, line, 1.0)
    else:
        shape.line.fill.background()
    return shape


def circle(slide, x, y, d, fill=SURFACE2, line=None, transparency=0):
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    set_fill(shape, fill, transparency)
    if line:
        set_line(shape, line, 1.0)
    else:
        shape.line.fill.background()
    return shape


def line(slide, x1, y1, x2, y2, color=MUTED2, width=1.5, dash=False, arrow=False):
    shp = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    set_line(shp, color, width)
    if dash:
        shp.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    if arrow:
        # python-pptx does not expose arrowheads consistently across versions.
        # A small triangle at the end keeps the visual editable and portable.
        ang = math.atan2(y2 - y1, x2 - x1)
        d = 0.12
        tri = slide.shapes.add_shape(
            MSO_SHAPE.ISOSCELES_TRIANGLE,
            Inches(x2 - d / 2), Inches(y2 - d / 2), Inches(d), Inches(d)
        )
        set_fill(tri, color)
        tri.line.fill.background()
        tri.rotation = math.degrees(ang) + 90
    return shp


def text(
    slide,
    value,
    x,
    y,
    w,
    h,
    size=16,
    color=WHITE,
    bold=False,
    font=FONT,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
    margin=0.02,
    italic=False,
    fit=False,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin)
    tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.text = str(value)
    p.alignment = align
    p.space_after = Pt(0)
    p.space_before = Pt(0)
    p.line_spacing = 1.02
    for run in p.runs:
        run.font.name = font
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.italic = italic
        run.font.color.rgb = C(color)
    if fit:
        try:
            tf.fit_text(font_family=font, max_size=Pt(size))
        except Exception:
            pass
    return box


def rich_text(slide, runs, x, y, w, h, size=16, align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = valign
    tf.margin_left = tf.margin_right = Inches(0.02)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    p = tf.paragraphs[0]
    p.alignment = align
    p.space_after = Pt(0)
    for spec in runs:
        r = p.add_run()
        r.text = spec[0]
        r.font.name = spec[4] if len(spec) > 4 else FONT
        r.font.size = Pt(spec[1] if len(spec) > 1 else size)
        r.font.bold = spec[2] if len(spec) > 2 else False
        r.font.color.rgb = C(spec[3] if len(spec) > 3 else WHITE)
    return box


def bullet_list(slide, items, x, y, w, h, size=15, color=WHITE, bullet_color=CYAN, gap=4):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.02)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_before = Pt(0)
        p.space_after = Pt(gap)
        p.line_spacing = 1.08
        r1 = p.add_run()
        r1.text = "•  "
        r1.font.name = FONT
        r1.font.size = Pt(size)
        r1.font.bold = True
        r1.font.color.rgb = C(bullet_color)
        r2 = p.add_run()
        r2.text = item
        r2.font.name = FONT
        r2.font.size = Pt(size)
        r2.font.color.rgb = C(color)
    return box


def label(slide, value, x, y, w, fill=PURPLE, color=WHITE, size=9.5):
    shape = rect(slide, x, y, w, 0.28, fill=fill, radius=True)
    text(slide, value.upper(), x, y + 0.005, w, 0.27, size=size, color=color, bold=True,
         align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    return shape


def footer(slide, number, source=None):
    line(slide, 0.55, 7.12, 12.78, 7.12, color=SURFACE2, width=0.8)
    text(slide, "VIRTUAL CELL CHALLENGE 2026", 0.58, 7.18, 3.2, 0.16, size=7.5, color=MUTED2, bold=True)
    if source:
        text(slide, source, 3.5, 7.16, 8.7, 0.2, size=7.2, color=MUTED2, align=PP_ALIGN.RIGHT)
    text(slide, f"{number:02d}", 12.35, 7.16, 0.4, 0.2, size=8, color=PURPLE2, bold=True,
         align=PP_ALIGN.RIGHT)


def blank_slide():
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = C(BG)
    return slide


def new_slide(section, title_value, subtitle=None, source=None, note=""):
    slide = blank_slide()
    n = len(prs.slides)
    rect(slide, 0, 0, W, 0.08, fill=PURPLE, radius=False)
    text(slide, section.upper(), 0.58, 0.30, 2.4, 0.24, size=9, color=CYAN, bold=True)
    text(slide, title_value, 0.58, 0.60, 12.1, 0.62, size=28, color=WHITE, bold=True)
    if subtitle:
        text(slide, subtitle, 0.60, 1.18, 11.9, 0.38, size=12.5, color=MUTED)
    footer(slide, n, source)
    SLIDES.append((n, title_value))
    NOTES.append((n, title_value, note))
    return slide


def add_note_to_slide(slide, note):
    if not note:
        return
    try:
        tf = slide.notes_slide.notes_text_frame
        tf.text = note
    except Exception:
        pass


def finish(slide, note):
    add_note_to_slide(slide, note)


def card(slide, x, y, w, h, title_value, body, accent=PURPLE, number=None, body_size=13):
    rect(slide, x, y, w, h, fill=SURFACE, line=SURFACE2)
    rect(slide, x, y, 0.06, h, fill=accent, radius=False)
    if number is not None:
        circle(slide, x + 0.20, y + 0.19, 0.38, fill=accent)
        text(slide, str(number), x + 0.20, y + 0.19, 0.38, 0.38, size=12, bold=True,
             align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
        tx = x + 0.70
        tw = w - 0.90
    else:
        tx = x + 0.26
        tw = w - 0.50
    text(slide, title_value, tx, y + 0.18, tw, 0.35, size=14, color=WHITE, bold=True)
    text(slide, body, x + 0.26, y + 0.65, w - 0.50, h - 0.82, size=body_size, color=MUTED)


def stat(slide, x, y, w, value, label_value, accent=CYAN, sub=None):
    rect(slide, x, y, w, 1.15, fill=SURFACE, line=SURFACE2)
    text(slide, value, x + 0.20, y + 0.14, w - 0.40, 0.42, size=24, color=accent, bold=True)
    text(slide, label_value, x + 0.20, y + 0.60, w - 0.40, 0.25, size=10.5, color=WHITE, bold=True)
    if sub:
        text(slide, sub, x + 0.20, y + 0.87, w - 0.40, 0.18, size=8.5, color=MUTED2)


def simple_table(slide, x, y, widths, row_h, headers, rows, header_fill=SURFACE2, font_size=10.5,
                 accents=None):
    xx = x
    for j, (h, w) in enumerate(zip(headers, widths)):
        rect(slide, xx, y, w, row_h, fill=header_fill, line=BG, radius=False)
        text(slide, h, xx + 0.08, y + 0.03, w - 0.16, row_h - 0.06, size=font_size, color=CYAN2,
             bold=True, valign=MSO_ANCHOR.MIDDLE)
        xx += w
    for i, row in enumerate(rows):
        xx = x
        fill = SURFACE if i % 2 == 0 else BG2
        for j, (v, w) in enumerate(zip(row, widths)):
            rect(slide, xx, y + row_h * (i + 1), w, row_h, fill=fill, line=BG, radius=False)
            col = accents[i] if accents and j == 0 else (WHITE if j == 0 else MUTED)
            text(slide, v, xx + 0.08, y + row_h * (i + 1) + 0.03, w - 0.16, row_h - 0.06,
                 size=font_size, color=col, bold=(j == 0), valign=MSO_ANCHOR.MIDDLE)
            xx += w


def small_cell(slide, x, y, d, color=PURPLE):
    circle(slide, x, y, d, fill=color, transparency=12)
    circle(slide, x + d * 0.34, y + d * 0.34, d * 0.32, fill=WHITE, transparency=15)
    for dx, dy in [(0.20, 0.18), (0.62, 0.28), (0.25, 0.66), (0.68, 0.66)]:
        circle(slide, x + d * dx, y + d * dy, d * 0.08, fill=CYAN, transparency=20)


def arrow_between(slide, x1, y1, x2, y2, color=CYAN):
    line(slide, x1, y1, x2, y2, color=color, width=2.0, arrow=True)


# 1 — title
slide = blank_slide()
rect(slide, 0, 0, W, H, fill=BG, radius=False)
for i, (x, y, d, col) in enumerate([
    (9.1, 0.55, 2.9, PURPLE), (10.65, 2.05, 1.5, CYAN), (8.35, 3.05, 2.2, PINK),
    (11.25, 4.10, 1.25, GREEN), (9.35, 5.15, 1.6, BLUE)
]):
    circle(slide, x, y, d, fill=col, transparency=42 if i else 25)
    circle(slide, x + d * 0.34, y + d * 0.34, d * 0.32, fill=BG, transparency=0)
label(slide, "Team discussion deck", 0.72, 0.68, 2.05, fill=PURPLE)
text(slide, "Virtual Cell\nChallenge 2026", 0.72, 1.28, 7.4, 1.70, size=38, color=WHITE, bold=True)
text(slide, "A zero-shot, single-cell perturbation prediction playbook", 0.76, 3.18, 7.6, 0.52,
     size=19, color=CYAN2, bold=True)
text(slide, "Biology  •  scoring  •  model strategy  •  three-person execution", 0.76, 3.86, 7.6, 0.38,
     size=13, color=MUTED)
rect(slide, 0.76, 4.70, 6.90, 1.25, fill=SURFACE, line=SURFACE2)
text(slide, "CURRENT TEAM SNAPSHOT", 1.02, 4.94, 2.20, 0.22, size=9.5, color=MUTED2, bold=True)
rich_text(slide, [("−0.0206", 26, True, RED), ("  overall  •  rank 173  •  Bayesian baseline v0", 12, False, WHITE)],
          1.02, 5.24, 6.35, 0.42)
text(slide, "Prepared 30 August 2026  |  Evidence snapshot + published H100 baseline", 0.76, 6.72, 8.1, 0.28,
     size=9, color=MUTED2)
text(slide, "01", 12.27, 6.98, 0.48, 0.24, size=9, color=PURPLE2, bold=True, align=PP_ALIGN.RIGHT)
SLIDES.append((1, "Virtual Cell Challenge 2026"))
note = (
    "Open with the decision problem, not the architecture. The current near-zero score is a normalized benchmark result, "
    "not an accuracy percentage. This deck explains the biological task, exact evaluation, and a three-person operating plan."
)
NOTES.append((1, "Virtual Cell Challenge 2026", note))
finish(slide, note)


# 2
note = "State the four strategic calls. The team should agree on these before discussing implementation details."
slide = new_slide("Executive decision", "The decision in 60 seconds",
                  "What we should optimize, what we should postpone, and what success means.",
                  "Evidence: official VCC contract + local audit", note)
card(slide, 0.62, 1.72, 3.00, 2.02, "Predict the perturbation delta first",
     "The six metrics reward target-specific mean direction and magnitude before they reward visual realism.", CYAN, 1, 12.5)
card(slide, 3.76, 1.72, 3.00, 2.02, "Generate cells around a trusted mean",
     "Use a count model or residual flow/diffusion to recover depth, variance, zeros, and responder heterogeneity.", PURPLE, 2, 12.5)
card(slide, 6.90, 1.72, 3.00, 2.02, "Use RL only behind a gate",
     "This is conditional distribution estimation, not sequential control. RL is optional verifier-guided post-training.", PINK, 3, 12.5)
card(slide, 10.04, 1.72, 2.68, 2.02, "Win on held-out contexts",
     "A leaderboard gain is evidence only when strict leave-one-cell-line-out evaluation also improves.", GREEN, 4, 12.0)
rect(slide, 0.62, 4.12, 12.10, 1.80, fill=BG2, line=PURPLE)
text(slide, "RECOMMENDED CORE", 0.93, 4.39, 1.75, 0.24, size=9.5, color=PURPLE2, bold=True)
text(slide, "context-conditioned deterministic effect predictor", 0.93, 4.79, 4.20, 0.43, size=18, bold=True)
text(slide, "+", 5.24, 4.78, 0.30, 0.42, size=19, color=CYAN, bold=True, align=PP_ALIGN.CENTER)
text(slide, "control-anchored count generator", 5.64, 4.79, 3.50, 0.43, size=18, bold=True)
text(slide, "+", 9.21, 4.78, 0.30, 0.42, size=19, color=CYAN, bold=True, align=PP_ALIGN.CENTER)
text(slide, "metric calibration + ensemble", 9.61, 4.79, 2.75, 0.43, size=16, bold=True)
text(slide, "Success = robust gain over the strongest statistical baseline on mean and worst held-out context.",
     0.93, 5.43, 11.20, 0.28, size=11.5, color=MUTED)
finish(slide, note)


# 3
note = "Use this map to orient a mixed biology/ML audience. The deck moves from assay to score to system to execution."
slide = new_slide("Executive decision", "Road map", "Five questions the team must answer together.", None, note)
chapters = [
    ("01", "What is measured?", "CRISPRi, Perturb-seq, raw UMI counts", CYAN),
    ("02", "What is predicted?", "A counterfactual 400-cell distribution", PURPLE),
    ("03", "How is it scored?", "Six metrics with distinct failure modes", PINK),
    ("04", "What should we build?", "Transfer model + calibrated count generator", GREEN),
    ("05", "How do three people deliver?", "Ownership, gates, timeline, submissions", YELLOW),
]
for i, (num, title_value, body, col) in enumerate(chapters):
    x = 0.64 + i * 2.48
    circle(slide, x + 0.78, 1.82, 0.78, fill=col, transparency=12)
    text(slide, num, x + 0.78, 1.82, 0.78, 0.78, size=16, bold=True,
         align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    if i < 4:
        arrow_between(slide, x + 1.60, 2.21, x + 2.29, 2.21, color=MUTED2)
    text(slide, title_value, x, 2.92, 2.28, 0.52, size=14, bold=True, align=PP_ALIGN.CENTER)
    text(slide, body, x, 3.53, 2.28, 0.88, size=11.5, color=MUTED, align=PP_ALIGN.CENTER)
rect(slide, 1.12, 5.20, 11.10, 0.78, fill=SURFACE, line=SURFACE2)
text(slide, "Discussion rule", 1.40, 5.43, 1.55, 0.25, size=11, color=CYAN, bold=True)
text(slide, "Every proposed model must identify which biological uncertainty and which metric failure it is designed to fix.",
     3.02, 5.39, 8.78, 0.32, size=12.5, color=WHITE)
finish(slide, note)


# 4
note = (
    "Define one condition as a question: given all non-targeting controls for an unseen context and one target gene, "
    "generate 400 post-CRISPRi cells over the fixed gene axis. There are 900 such questions in a round."
)
slide = new_slide("The task", "The challenge in one sentence",
                  "Infer a missing intervention distribution from an observed control population.",
                  "Source: VCC 2026 Data; VCC CLI Guide", note)
rect(slide, 0.65, 1.75, 3.10, 3.80, fill=SURFACE, line=CYAN)
label(slide, "Observed", 0.92, 2.02, 1.05, fill=CYAN, color=BG)
for i in range(8):
    small_cell(slide, 0.95 + (i % 4) * 0.58, 2.58 + (i // 4) * 0.62, 0.43, [CYAN, BLUE, PURPLE][i % 3])
text(slide, "18,400 non-targeting\ncontrol cells", 0.95, 3.98, 2.35, 0.72, size=17, bold=True)
text(slide, "Anonymous context c", 0.95, 4.80, 2.35, 0.28, size=11, color=MUTED)
rect(slide, 4.28, 1.75, 2.20, 3.80, fill=BG2, line=PURPLE)
label(slide, "Condition", 4.58, 2.02, 1.05, fill=PURPLE)
circle(slide, 4.88, 2.75, 1.00, fill=PURPLE, transparency=15)
text(slide, "p", 4.88, 2.75, 1.00, 1.00, size=28, bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
text(slide, "CRISPRi target\ngene", 4.62, 4.05, 1.54, 0.68, size=16, bold=True, align=PP_ALIGN.CENTER)
arrow_between(slide, 3.78, 3.64, 4.18, 3.64, color=CYAN)
arrow_between(slide, 6.58, 3.64, 7.04, 3.64, color=CYAN)
rect(slide, 7.12, 1.75, 5.58, 3.80, fill=SURFACE, line=GREEN)
label(slide, "Counterfactual", 7.42, 2.02, 1.38, fill=GREEN, color=BG)
for i in range(12):
    small_cell(slide, 7.48 + (i % 6) * 0.60, 2.64 + (i // 6) * 0.62, 0.42,
               [GREEN, PURPLE, PINK, BLUE][i % 4])
text(slide, "400 predicted cells × 18,533 genes", 7.46, 4.10, 4.55, 0.42, size=17, bold=True)
text(slide, "Raw nonnegative whole counts • no controls in upload", 7.46, 4.70, 4.55, 0.35, size=11, color=MUTED)
text(slide, "(D⁰c, p)  →  X̂c,p  ≈  P(X | do[CRISPRi(p)], c)", 1.18, 5.96, 10.95, 0.44,
     size=20, color=CYAN2, bold=True, font=MONO, align=PP_ALIGN.CENTER)
finish(slide, note)


# 5
note = (
    "CRISPRi uses a catalytically inactive Cas9 recruited to a target locus, commonly with a repressor domain, to suppress transcription. "
    "It is a knockdown rather than a DNA-cutting knockout. The challenge reports greater than 80% on-target knockdown for scored targets."
)
slide = new_slide("Biology primer", "What CRISPR interference changes",
                  "A guide-directed transcriptional brake creates a causal perturbation without cutting DNA.",
                  "Sources: Qi et al., Cell (2013); VCC 2026 Data", note)
text(slide, "DNA", 0.70, 1.78, 0.55, 0.26, size=11, color=MUTED, bold=True)
line(slide, 1.22, 1.92, 7.30, 1.92, color=BLUE, width=3)
for i in range(14):
    line(slide, 1.35 + i * 0.41, 1.78, 1.35 + i * 0.41, 2.06, color=CYAN2, width=1.1)
rect(slide, 3.15, 1.60, 1.50, 0.62, fill=PURPLE, line=PURPLE2)
text(slide, "dCas9–repressor", 3.21, 1.73, 1.38, 0.26, size=10.5, bold=True, align=PP_ALIGN.CENTER)
line(slide, 2.20, 2.70, 4.02, 2.04, color=PINK, width=2.3)
text(slide, "sgRNA", 1.44, 2.63, 0.72, 0.28, size=11, color=PINK2, bold=True)
arrow_between(slide, 4.78, 1.92, 6.18, 1.92, color=RED)
line(slide, 5.08, 1.48, 5.08, 2.36, color=RED, width=4)
text(slide, "transcription blocked", 5.30, 2.34, 2.05, 0.30, size=11.5, color=RED, bold=True)
rect(slide, 0.70, 3.20, 6.70, 2.34, fill=SURFACE, line=SURFACE2)
text(slide, "Mechanistic consequences", 0.98, 3.46, 2.78, 0.34, size=16, bold=True)
bullet_list(slide, [
    "Target RNA decreases; the target transcript itself is excluded from scoring.",
    "Downstream genes respond through regulatory, signaling, metabolic, and stress networks.",
    "The same target can induce different responses in different cell contexts.",
    "A knockdown is not necessarily equivalent to a complete genetic knockout."
], 0.98, 3.94, 5.94, 1.34, size=12.5, bullet_color=PURPLE2, gap=3)
rect(slide, 7.78, 1.62, 4.52, 3.92, fill=BG2, line=CYAN)
text(slide, ">80%", 8.18, 2.02, 3.72, 0.72, size=38, color=CYAN, bold=True, align=PP_ALIGN.CENTER)
text(slide, "reported on-target knockdown", 8.18, 2.82, 3.72, 0.36, size=14, bold=True, align=PP_ALIGN.CENTER)
line(slide, 8.32, 3.48, 11.75, 3.48, color=SURFACE2, width=1)
text(slide, "Model implication", 8.18, 3.74, 3.72, 0.32, size=12, color=PURPLE2, bold=True, align=PP_ALIGN.CENTER)
text(slide, "Do not learn only ‘target goes down.’\nLearn the trans-effect network and its context modulation.",
     8.18, 4.20, 3.72, 0.80, size=14, color=WHITE, align=PP_ALIGN.CENTER)
finish(slide, note)


# 6
note = (
    "Perturb-seq links a guide identity to a whole-transcriptome readout in each cell. Pooling provides scale; single-cell measurements "
    "reveal heterogeneous responses that pseudobulk alone cannot represent."
)
slide = new_slide("Biology primer", "What Perturb-seq measures",
                  "Guide identity and a transcriptome are read together for each cell.",
                  "Source: Dixit et al., Cell (2016)", note)
stages = [
    ("1", "Pooled guides", "Many sgRNAs enter one experiment", PURPLE),
    ("2", "CRISPRi cells", "Each cell carries a perturbation identity", PINK),
    ("3", "10x capture", "RNA molecules receive cell and UMI barcodes", CYAN),
    ("4", "Count matrix", "One sparse transcriptome per cell", GREEN),
]
for i, (num, ttl, body, col) in enumerate(stages):
    x = 0.68 + i * 3.06
    rect(slide, x, 1.82, 2.60, 2.86, fill=SURFACE, line=col)
    circle(slide, x + 0.18, 2.04, 0.46, fill=col)
    text(slide, num, x + 0.18, 2.04, 0.46, 0.46, size=13, bold=True,
         align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    if i == 0:
        for j in range(5):
            line(slide, x + 0.55 + j * 0.30, 2.89, x + 0.82 + j * 0.30, 3.49,
                 color=[PURPLE2, CYAN2, PINK2, GREEN2, YELLOW2][j], width=2.2)
    elif i == 1:
        for j in range(4):
            small_cell(slide, x + 0.44 + (j % 2) * 0.74, 2.80 + (j // 2) * 0.63, 0.52,
                       [PURPLE, PINK, BLUE, GREEN][j])
    elif i == 2:
        circle(slide, x + 0.65, 2.74, 1.18, fill=CYAN, transparency=20)
        for j in range(6):
            line(slide, x + 0.92, 2.98 + j * 0.10, x + 1.60, 2.85 + j * 0.16,
                 color=CYAN2, width=1.0)
    else:
        for r in range(5):
            for c in range(7):
                colr = GREEN if (r * 7 + c) % 4 == 0 else SURFACE2
                rect(slide, x + 0.36 + c * 0.23, 2.73 + r * 0.20, 0.18, 0.14, fill=colr, radius=False)
    text(slide, ttl, x + 0.22, 3.89, 2.16, 0.34, size=14, bold=True, align=PP_ALIGN.CENTER)
    text(slide, body, x + 0.22, 4.26, 2.16, 0.52, size=10.7, color=MUTED, align=PP_ALIGN.CENTER)
    if i < 3:
        arrow_between(slide, x + 2.68, 3.24, x + 2.94, 3.24, color=MUTED2)
rect(slide, 1.10, 5.20, 11.10, 0.82, fill=BG2, line=SURFACE2)
text(slide, "The readout is not one phenotype.", 1.38, 5.43, 3.36, 0.30, size=13, color=CYAN2, bold=True)
text(slide, "It is a 18,533-dimensional molecular phenotype with technical sampling noise and biological heterogeneity.",
     4.67, 5.41, 7.06, 0.34, size=12.5)
finish(slide, note)


# 7
note = (
    "Explain rows, columns, counts, and zeros. A zero can reflect true absence, low abundance, or finite UMI sampling. "
    "The matrix is sparse in storage even though this dataset has substantial per-cell coverage."
)
slide = new_slide("Biology primer", "How to read a single-cell count matrix",
                  "Rows are cells; columns are genes; entries are captured molecule counts.",
                  "Evidence: local audit of controls.zip", note)
# Matrix
genes = ["TSPAN6", "TNMD", "DPM1", "SCYL3", "…", "ZRANB3"]
for j, g in enumerate(genes):
    text(slide, g, 2.02 + j * 0.82, 1.73, 0.78, 0.28, size=9, color=CYAN2, bold=True, align=PP_ALIGN.CENTER)
vals = [
    [0, 0, 2, 1, 0, 0], [1, 0, 0, 4, 0, 1], [0, 0, 1, 0, 3, 0],
    [2, 0, 0, 1, 0, 0], [0, 0, 5, 0, 1, 0], [1, 0, 0, 2, 0, 0],
]
for i, row in enumerate(vals):
    text(slide, f"cell {i+1}", 0.86, 2.12 + i * 0.52, 0.85, 0.30, size=10, color=MUTED, align=PP_ALIGN.RIGHT)
    for j, v in enumerate(row):
        intensity = [SURFACE, "1B3451", "215C75", "1D8390", "1CA5A9", CYAN][min(v, 5)]
        rect(slide, 2.08 + j * 0.82, 2.03 + i * 0.52, 0.66, 0.40, fill=intensity, line=BG, radius=False)
        text(slide, str(v), 2.08 + j * 0.82, 2.04 + i * 0.52, 0.66, 0.38, size=10,
             color=WHITE if v else MUTED2, bold=bool(v), align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
line(slide, 7.45, 1.72, 7.45, 5.48, color=SURFACE2, width=1.2)
card(slide, 7.82, 1.75, 4.70, 1.05, "Raw UMI-like counts",
     "Submit whole, finite, nonnegative values—not log1p, CPM, z-scores, or latent vectors.", CYAN, body_size=11.5)
card(slide, 7.82, 2.97, 4.70, 1.05, "Zeros are ambiguous",
     "A zero may be biological or a consequence of finite molecular capture. Avoid treating every zero as dropout to impute.", PURPLE, body_size=11.0)
card(slide, 7.82, 4.19, 4.70, 1.05, "Sparsity is operational",
     "Dense submission shape exceeds the stored-entry cap; CSR/CSC and removal of explicit zeros are mandatory.", GREEN, body_size=11.0)
rect(slide, 0.86, 5.72, 11.66, 0.54, fill=BG2, line=SURFACE2)
text(slide, "Model in a normalized latent space if useful—then reconstruct and validate raw-count space before packaging.",
     1.10, 5.87, 11.16, 0.25, size=11.5, color=YELLOW2, bold=True, align=PP_ALIGN.CENTER)
finish(slide, note)


# 8
note = (
    "There is no one-to-one pairing between predicted and measured cells. Four hundred cells determine group means, variance, "
    "zero fractions, and Wilcoxon power. Repeating one centroid 400 times can distort differential-expression calls."
)
slide = new_slide("Biology primer", "Why 400 cells per perturbation matter",
                  "The evaluator compares distributions and group summaries—not matched cell identities.",
                  "Sources: VCC metric specification; VCC CLI Guide", note)
card(slide, 0.70, 1.76, 3.70, 3.52, "What 400 cells estimate",
     "• a stable pseudobulk centroid\n\n• gene-wise dispersion and zero fraction\n\n• responder versus non-responder mixtures\n\n• statistical power for Wilcoxon DE calls",
     CYAN, body_size=13)
rect(slide, 4.72, 1.76, 3.72, 3.52, fill=SURFACE, line=PURPLE)
text(slide, "Two outputs can share a mean…", 5.02, 2.04, 3.12, 0.34, size=14, bold=True, align=PP_ALIGN.CENTER)
for j in range(16):
    x = 5.16 + (j % 8) * 0.34
    y = 2.68 + (j // 8) * 0.46
    circle(slide, x, y, 0.18, fill=PURPLE, transparency=10)
for j in range(16):
    theta = 2 * math.pi * j / 16
    x = 6.32 + 0.92 * math.cos(theta)
    y = 4.22 + 0.46 * math.sin(theta)
    circle(slide, x, y, 0.16, fill=PINK if j % 3 else CYAN, transparency=10)
text(slide, "…but imply different DE power.", 5.02, 4.77, 3.12, 0.30, size=12, color=MUTED, align=PP_ALIGN.CENTER)
card(slide, 8.76, 1.76, 3.86, 1.50, "Mode collapse",
     "Duplicating an identical cell makes variance unrealistically small and can over-call significance.", RED, body_size=11.5)
card(slide, 8.76, 3.45, 3.86, 1.50, "Excess noise",
     "Over-dispersed samples can hide true DE genes and damage direction reach and Jaccard.", YELLOW, body_size=11.5)
text(slide, "Mean accuracy is necessary; calibrated variability is the second-stage requirement.",
     1.00, 5.80, 11.34, 0.40, size=14, color=GREEN2, bold=True, align=PP_ALIGN.CENTER)
finish(slide, note)


# 9
note = (
    "The anonymous label is not the context representation. The control population contains basal state, cell-state composition, "
    "depth, covariance, and technical noise. Preserve A/B/C labels exactly; swapping labels can look like a weak model rather than a file error."
)
slide = new_slide("Biology primer", "Control cells define the context",
                  "A context is an empirical distribution—not merely the letter A, B, or C.",
                  "Sources: VCC 2026 Data; VCC CLI Guide", note)
small_cell(slide, 0.86, 2.08, 1.42, CYAN)
text(slide, "NTC population", 0.72, 3.67, 1.72, 0.32, size=13, bold=True, align=PP_ALIGN.CENTER)
arrow_between(slide, 2.55, 2.79, 3.22, 2.79, color=CYAN)
rect(slide, 3.30, 1.78, 4.10, 2.10, fill=SURFACE, line=CYAN)
text(slide, "Context encoder", 3.64, 2.05, 3.42, 0.38, size=18, bold=True, align=PP_ALIGN.CENTER)
features = ["pseudobulk", "cell-state mixture", "depth", "dispersion", "gene modules", "covariance"]
for i, f in enumerate(features):
    label(slide, f, 3.58 + (i % 3) * 1.15, 2.63 + (i // 3) * 0.50, 1.02,
          fill=[CYAN, PURPLE, BLUE, GREEN, PINK, YELLOW][i], color=BG if i in [0,3,5] else WHITE, size=7.8)
arrow_between(slide, 7.53, 2.79, 8.18, 2.79, color=CYAN)
circle(slide, 8.32, 2.12, 1.34, fill=PURPLE, transparency=10)
text(slide, "zc", 8.32, 2.12, 1.34, 1.34, size=26, bold=True, font=MONO,
     align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
arrow_between(slide, 9.88, 2.79, 10.42, 2.79, color=CYAN)
rect(slide, 10.56, 1.95, 1.94, 1.68, fill=BG2, line=PURPLE)
text(slide, "modulate\ntarget effect", 10.76, 2.30, 1.54, 0.72, size=15, bold=True, align=PP_ALIGN.CENTER)
rect(slide, 0.80, 4.48, 11.72, 1.28, fill=SURFACE, line=SURFACE2)
text(slide, "Critical implementation rule", 1.08, 4.77, 2.42, 0.30, size=12, color=RED, bold=True)
text(slide, "Attach the original context label at ingestion and never infer, reorder, or relabel it downstream.",
     3.60, 4.72, 8.48, 0.38, size=14, bold=True)
text(slide, "The package validator checks label presence, but it cannot detect a biologically swapped A/B mapping.",
     3.60, 5.20, 8.48, 0.28, size=10.5, color=MUTED)
finish(slide, note)


# 10
note = (
    "For each target we observe controls but not the matched counterfactual in challenge contexts. Model the response as a delta from "
    "the context-specific control profile. Separating baseline and effect improves transfer and calibration."
)
slide = new_slide("Problem formulation", "The causal quantity of interest",
                  "Predict the response to an intervention, not the identity of a cell line.",
                  "Interpretation: team modeling framework", note)
text(slide, "Observed", 0.82, 1.78, 1.20, 0.30, size=11, color=CYAN, bold=True)
rect(slide, 0.82, 2.19, 3.20, 2.36, fill=SURFACE, line=CYAN)
text(slide, "μ⁰c", 1.50, 2.66, 1.84, 0.82, size=36, color=CYAN2, bold=True, font=MONO, align=PP_ALIGN.CENTER)
text(slide, "control-state mean", 1.30, 3.62, 2.24, 0.28, size=12, color=MUTED, align=PP_ALIGN.CENTER)
text(slide, "+", 4.26, 2.88, 0.66, 0.64, size=30, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
text(slide, "Missing effect", 5.02, 1.78, 1.60, 0.30, size=11, color=PURPLE2, bold=True)
rect(slide, 5.02, 2.19, 3.20, 2.36, fill=SURFACE, line=PURPLE)
text(slide, "Δc,p", 5.68, 2.66, 1.88, 0.82, size=36, color=PURPLE2, bold=True, font=MONO, align=PP_ALIGN.CENTER)
text(slide, "context × target response", 5.40, 3.62, 2.42, 0.28, size=12, color=MUTED, align=PP_ALIGN.CENTER)
text(slide, "=", 8.42, 2.88, 0.66, 0.64, size=30, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
text(slide, "Counterfactual", 9.18, 1.78, 1.90, 0.30, size=11, color=GREEN, bold=True)
rect(slide, 9.18, 2.19, 3.20, 2.36, fill=SURFACE, line=GREEN)
text(slide, "μ¹c,p", 9.70, 2.66, 2.16, 0.82, size=34, color=GREEN2, bold=True, font=MONO, align=PP_ALIGN.CENTER)
text(slide, "post-CRISPRi mean", 9.64, 3.62, 2.26, 0.28, size=12, color=MUTED, align=PP_ALIGN.CENTER)
rect(slide, 1.34, 5.05, 10.62, 0.84, fill=BG2, line=SURFACE2)
text(slide, "Then generate individual cells around μ¹c,p with context-matched depth, dispersion, zeros, and state mixtures.",
     1.68, 5.31, 9.94, 0.32, size=13, color=YELLOW2, bold=True, align=PP_ALIGN.CENTER)
finish(slide, note)


# 11
note = (
    "Final evaluation changes both cell contexts and target panel. In addition, public training data may differ in assay, time point, "
    "perturbation modality, and sequencing depth. Use target representations and context representations that extrapolate."
)
slide = new_slide("Problem formulation", "Zero-shot means multiple extrapolations",
                  "The hidden test does not preserve the identities that a one-hot model can memorize.",
                  "Sources: VCC 2026 Data; VCC CLI Guide", note)
axes = [
    ("Context shift", "A / B / C  →  D / E / F", "Three entirely different anonymous cell lines", CYAN),
    ("Target shift", "validation panel  →  final panel", "A different set of 300 CRISPRi targets", PURPLE),
    ("Domain shift", "public assays  →  Arc 10x Flex", "Platform, depth, guides, timing, and batch", PINK),
    ("Distribution shift", "mean response  →  400 cells", "Variance, zeros, state mixtures, and responders", GREEN),
]
for i, (ttl, path, body, col) in enumerate(axes):
    y = 1.72 + i * 1.16
    rect(slide, 0.80, y, 2.06, 0.88, fill=col, transparency=8)
    text(slide, ttl, 0.97, y + 0.23, 1.72, 0.32, size=14, color=BG if col in [CYAN, GREEN, YELLOW] else WHITE,
         bold=True, align=PP_ALIGN.CENTER)
    rect(slide, 3.08, y, 3.64, 0.88, fill=BG2, line=col)
    text(slide, path, 3.30, y + 0.22, 3.20, 0.34, size=13, bold=True, font=MONO, align=PP_ALIGN.CENTER)
    rect(slide, 6.94, y, 5.50, 0.88, fill=SURFACE, line=SURFACE2)
    text(slide, body, 7.20, y + 0.22, 4.98, 0.38, size=12.5, color=MUTED, align=PP_ALIGN.CENTER)
text(slide, "Design consequence", 0.82, 6.44, 1.42, 0.24, size=10, color=YELLOW, bold=True)
text(slide, "No context-ID lookup table. No target-ID-only embedding. No validation-leaderboard tuning as a training objective.",
     2.38, 6.38, 9.82, 0.34, size=13, color=WHITE, bold=True)
finish(slide, note)


# 12
note = (
    "Keep development and final evaluation conceptually separate. Final controls and a different target panel arrive October 22. "
    "The architecture and decision rule should be frozen before that date."
)
slide = new_slide("Problem formulation", "Validation is not the final test",
                  "The final phase replaces both the contexts and perturbation panel.",
                  "Sources: VCC 2026 Data; VCC CLI Guide; challenge portal", note)
rect(slide, 0.72, 1.72, 5.74, 3.84, fill=SURFACE, line=CYAN)
label(slide, "Development round", 1.02, 1.98, 1.52, fill=CYAN, color=BG)
text(slide, "A  •  B  •  C", 1.02, 2.60, 4.80, 0.55, size=28, color=CYAN2, bold=True, align=PP_ALIGN.CENTER)
text(slide, "3 anonymous validation cell lines", 1.02, 3.22, 4.80, 0.30, size=13, align=PP_ALIGN.CENTER)
text(slide, "300 shared validation targets", 1.02, 3.66, 4.80, 0.30, size=13, color=MUTED, align=PP_ALIGN.CENTER)
text(slide, "Live leaderboard • development evidence", 1.02, 4.42, 4.80, 0.30, size=11.5, color=YELLOW2, align=PP_ALIGN.CENTER)
rect(slide, 6.86, 1.72, 5.74, 3.84, fill=SURFACE, line=PURPLE)
label(slide, "Prize round", 7.16, 1.98, 1.13, fill=PURPLE)
text(slide, "D  •  E  •  F", 7.16, 2.60, 4.80, 0.55, size=28, color=PURPLE2, bold=True, align=PP_ALIGN.CENTER)
text(slide, "3 different final cell lines", 7.16, 3.22, 4.80, 0.30, size=13, align=PP_ALIGN.CENTER)
text(slide, "A different panel of 300 targets", 7.16, 3.66, 4.80, 0.30, size=13, color=MUTED, align=PP_ALIGN.CENTER)
text(slide, "Controls released 22 Oct • submissions close 5 Nov", 7.16, 4.42, 4.80, 0.30,
     size=11.5, color=YELLOW2, align=PP_ALIGN.CENTER)
arrow_between(slide, 6.47, 3.62, 6.77, 3.62, color=MUTED2)
rect(slide, 1.52, 5.92, 10.30, 0.55, fill=BG2, line=RED)
text(slide, "Freeze architecture, evaluation, and ensemble rules before 22 October; adapt only through the released controls.",
     1.82, 6.06, 9.70, 0.28, size=12.3, color=RED, bold=True, align=PP_ALIGN.CENTER)
finish(slide, note)


# 13
note = (
    "The upload contains exactly 360,000 predicted cells: 300 targets times 400 cells times 3 contexts. All 18,533 genes are present. "
    "One sparse .vcc file is required and controls must not be included."
)
slide = new_slide("Data contract", "The exact prediction contract",
                  "A complete round is large in logical shape but manageable when stored sparsely.",
                  "Source: VCC CLI Guide — Submission requirements (2026)", note)
stat(slide, 0.68, 1.70, 2.72, "300", "targets / context", PURPLE, "same panel across three contexts")
stat(slide, 3.58, 1.70, 2.72, "400", "cells / target", CYAN, "exactly—not a minimum")
stat(slide, 6.48, 1.70, 2.72, "3", "contexts / round", GREEN, "A/B/C or D/E/F")
stat(slide, 9.38, 1.70, 2.72, "18,533", "genes / cell", PINK, "official fixed axis")
text(slide, "300 × 400 × 3 =", 1.26, 3.30, 3.25, 0.55, size=22, color=MUTED, bold=True, align=PP_ALIGN.RIGHT)
text(slide, "360,000 cells", 4.70, 3.27, 3.20, 0.62, size=30, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
text(slide, "× 18,533 genes", 8.04, 3.30, 3.30, 0.55, size=22, color=MUTED, bold=True)
rect(slide, 0.92, 4.36, 11.48, 1.50, fill=SURFACE, line=SURFACE2)
contracts = [
    ("RAW", "whole nonnegative counts", CYAN), ("SPARSE", "CSR / CSC; remove explicit zeros", PURPLE),
    ("LABELS", "exact context + target symbols", PINK), ("ONE FILE", "all three contexts; no NTC rows", GREEN),
]
for i, (ttl, body, col) in enumerate(contracts):
    x = 1.18 + i * 2.82
    label(slide, ttl, x, 4.68, 0.88, fill=col, color=BG if col in [CYAN, GREEN] else WHITE)
    text(slide, body, x, 5.13, 2.30, 0.46, size=10.5, color=MUTED)
text(slide, "Dense logical positions: 6,671,880,000  >  stored-entry cap: 4,750,000,000",
     1.34, 6.25, 10.66, 0.34, size=12.5, color=YELLOW2, bold=True, font=MONO, align=PP_ALIGN.CENTER)
finish(slide, note)


# 14
note = (
    "Summarize what was actually downloaded and audited. Each context contains 18,400 NTC cells: 46 guide IDs times 400 cells. "
    "All numeric and axis checks passed. Context B is somewhat sparser and shallower than A/C."
)
slide = new_slide("Data contract", "Downloaded controls: measured facts",
                  "The local archive passed structural, numerical, and metadata checks.",
                  "Evidence: artifacts/control_audit.json", note)
simple_table(slide, 0.72, 1.74, [1.25, 1.45, 1.62, 1.62, 1.62, 1.58, 1.62], 0.58,
             ["Context", "Cells", "Density", "Mean genes/cell", "Mean UMI/cell", "Median UMI", "All-zero genes"],
             [
                 ["A", "18,400", "32.23%", "5,973", "21,134", "20,109", "2,398"],
                 ["B", "18,400", "29.78%", "5,519", "19,990", "19,946", "2,853"],
                 ["C", "18,400", "31.73%", "5,881", "21,157", "20,034", "2,317"],
             ], font_size=10.0, accents=[CYAN, PURPLE, PINK])
stat(slide, 0.72, 4.34, 2.70, "55,200", "total control cells", CYAN)
stat(slide, 3.58, 4.34, 2.70, "46 × 400", "NTC design / context", PURPLE)
stat(slide, 6.44, 4.34, 2.70, "319.7M", "stored nonzeros", GREEN)
stat(slide, 9.30, 4.34, 2.70, "662 MB", "controls.zip", PINK)
rect(slide, 0.72, 5.78, 11.28, 0.56, fill=BG2, line=GREEN)
text(slide, "PASS: finite • nonnegative • integer-valued • canonical CSR • identical gene order • 300 unique targets",
     0.94, 5.92, 10.84, 0.28, size=11.8, color=GREEN2, bold=True, align=PP_ALIGN.CENTER)
finish(slide, note)


# 15
note = (
    "The control centroids are not interchangeable. Pairwise cosine similarities in a smoke subset range from 0.721 to 0.808. "
    "This is evidence that context conditioning matters, not an attempt to identify the anonymous lines."
)
slide = new_slide("Data contract", "The three contexts are genuinely different",
                  "Control-centroid similarity is high enough to share structure, but low enough to require context modulation.",
                  "Evidence: artifacts/h100_smoke_result.json", note)
labels_hm = ["A", "B", "C"]
mat = [[1.000, 0.778, 0.721], [0.778, 1.000, 0.808], [0.721, 0.808, 1.000]]
text(slide, "Cosine similarity of log-CP10K control centroids", 0.86, 1.76, 5.60, 0.34, size=15, bold=True)
for j, lab in enumerate(labels_hm):
    text(slide, lab, 2.06 + j * 1.02, 2.18, 0.82, 0.30, size=12, color=CYAN2, bold=True, align=PP_ALIGN.CENTER)
for i, lab in enumerate(labels_hm):
    text(slide, lab, 1.20, 2.65 + i * 0.90, 0.52, 0.32, size=12, color=CYAN2, bold=True, align=PP_ALIGN.CENTER)
    for j in range(3):
        v = mat[i][j]
        col = GREEN if v > 0.95 else (CYAN if v > 0.80 else (PURPLE if v > 0.75 else PINK))
        rect(slide, 2.06 + j * 1.02, 2.50 + i * 0.90, 0.82, 0.70, fill=col, transparency=int((1 - v) * 55), radius=False)
        text(slide, f"{v:.3f}", 2.06 + j * 1.02, 2.50 + i * 0.90, 0.82, 0.70, size=11,
             color=BG if col in [GREEN, CYAN] else WHITE, bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
rect(slide, 6.50, 1.74, 5.80, 3.58, fill=SURFACE, line=SURFACE2)
text(slide, "Interpretation", 6.84, 2.02, 2.00, 0.34, size=16, color=PURPLE2, bold=True)
bullet_list(slide, [
    "Shared programs make cross-context transfer plausible.",
    "Distinct basal states make a single global response insufficient.",
    "Context identity must come from the control distribution, not a memorized letter.",
    "Do not infer the hidden biological identity of A, B, or C."
], 6.84, 2.52, 5.04, 2.26, size=12.2, bullet_color=CYAN)
rect(slide, 0.86, 5.73, 11.44, 0.64, fill=BG2, line=PURPLE)
text(slide, "Design target: learn Δglobal(p) + Δinteraction(context, p), with shrinkage when evidence is weak.",
     1.12, 5.91, 10.92, 0.30, size=13, color=PURPLE2, bold=True, font=MONO, align=PP_ALIGN.CENTER)
finish(slide, note)


# 16
note = (
    "Report the actual infrastructure result. Job 792580 completed on an H100 80GB. It loaded real controls, ran BF16 forward/backward/update, "
    "built a full sparse 360,000-cell artifact, passed the official dry run, and packaged a .vcc file."
)
slide = new_slide("H100 readiness", "The end-to-end H100 smoke test passed",
                  "Data → CUDA training path → sparse H5AD → official validation → .vcc packaging.",
                  "Evidence: job 792580; reports/VCC_2026_H100_smoke_runbook_bilingual.md", note)
stages = [
    ("01", "Data audit", "55,200 controls", CYAN),
    ("02", "BF16 train", "9.54M parameters", PURPLE),
    ("03", "Full artifact", "360,000 cells", PINK),
    ("04", "Official prep", "targets verified", GREEN),
    ("05", "Package", "16.46 MB .vcc", YELLOW),
]
for i, (num, ttl, body, col) in enumerate(stages):
    x = 0.62 + i * 2.50
    circle(slide, x + 0.73, 1.78, 0.70, fill=col, transparency=8)
    text(slide, "✓", x + 0.73, 1.78, 0.70, 0.70, size=20, color=BG if col in [CYAN, GREEN, YELLOW] else WHITE,
         bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    text(slide, ttl, x, 2.72, 2.16, 0.34, size=14, bold=True, align=PP_ALIGN.CENTER)
    text(slide, body, x, 3.15, 2.16, 0.28, size=10.5, color=MUTED, align=PP_ALIGN.CENTER)
    if i < 4:
        arrow_between(slide, x + 1.88, 2.14, x + 2.25, 2.14, color=MUTED2)
simple_table(slide, 1.12, 4.18, [2.25, 2.10, 2.40, 2.48, 2.18], 0.50,
             ["GPU", "Elapsed", "Peak RSS", "Training loss", "Exit"],
             [["H100 80GB", "1m 31s", "≈4.91 GiB", "0.3466 → 0.2165", "COMPLETED"]],
             font_size=10.5, accents=[GREEN])
text(slide, "The pipeline is competition-ready; the prediction model is not yet scientifically ready.",
     1.18, 5.82, 10.98, 0.40, size=15, color=YELLOW2, bold=True, align=PP_ALIGN.CENTER)
finish(slide, note)


# 17
note = (
    "Be explicit: the smoke output repeats a control-derived null population across all targets. It is named DO_NOT_SUBMIT. "
    "It proves mechanics, not target-specific biological prediction or leaderboard performance."
)
slide = new_slide("H100 readiness", "What the smoke test did not prove",
                  "Passing infrastructure gates is not evidence of perturbation biology.",
                  "Evidence: local smoke artifact metadata", note)
rect(slide, 0.72, 1.70, 5.48, 3.90, fill=SURFACE, line=GREEN)
label(slide, "Proven", 1.02, 1.98, 0.88, fill=GREEN, color=BG)
bullet_list(slide, [
    "H100 allocation and BF16 tensor path work.",
    "Real controls can be loaded without densifying the full matrix.",
    "Sparse 360,000 × 18,533 output can be written.",
    "Official target, context, count, and gene checks pass.",
    "The .vcc packaging workflow is reproducible."
], 1.04, 2.56, 4.62, 2.48, size=12.5, bullet_color=GREEN)
rect(slide, 6.56, 1.70, 5.98, 3.90, fill=SURFACE, line=RED)
label(slide, "Not proven", 6.86, 1.98, 1.08, fill=RED)
bullet_list(slide, [
    "No target-specific response was learned.",
    "No cross-context generalization was evaluated.",
    "No six-metric gain was established.",
    "No public perturbation data were used.",
    "The generated null artifact must not be submitted."
], 6.88, 2.56, 5.10, 2.48, size=12.5, bullet_color=RED)
rect(slide, 1.15, 5.98, 11.02, 0.50, fill=RED, transparency=8, line=RED)
text(slide, "DO_NOT_SUBMIT = an intentional safety label, not a candidate model name.",
     1.42, 6.10, 10.48, 0.26, size=12.5, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
finish(slide, note)


# 18
note = (
    "The normalized score uses two anchors per context and metric: zero is a context-mean perturbation-response baseline; one is split-half "
    "experimental reproducibility. Negative means worse than baseline. The current −0.0206 is not −2.06% accuracy."
)
slide = new_slide("Evaluation", "The score is a ruler—not a percentage",
                  "Two biological anchors make six different metrics comparable.",
                  "Sources: VCC Evaluation; VCC CLI Guide; cell-eval2 metric specification", note)
# Gauge
line(slide, 1.02, 3.08, 12.15, 3.08, color=MUTED2, width=4)
for xv, val, col, ttl in [(4.35, "−0.0206", RED, "current team"), (5.28, "0", YELLOW, "mean-response baseline"), (9.44, "1", GREEN, "split-half replicate")]:
    circle(slide, xv, 2.70, 0.76, fill=col)
    text(slide, val, xv - 0.32, 1.92, 1.40, 0.42, size=18 if val != "−0.0206" else 16,
         color=col, bold=True, align=PP_ALIGN.CENTER)
    text(slide, ttl, xv - 0.62, 3.68, 2.00, 0.58, size=11.5, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
text(slide, "worse than baseline", 0.96, 4.62, 3.14, 0.32, size=11, color=MUTED, align=PP_ALIGN.CENTER)
text(slide, "meaningful model progress", 4.16, 4.62, 4.28, 0.32, size=11, color=MUTED, align=PP_ALIGN.CENTER)
text(slide, "experimental reproducibility", 8.50, 4.62, 3.82, 0.32, size=11, color=MUTED, align=PP_ALIGN.CENTER)
rect(slide, 1.12, 5.34, 11.08, 0.82, fill=BG2, line=SURFACE2)
text(slide, "Overall = unweighted mean of 18 context × metric normalized cells", 1.42, 5.53, 5.98, 0.32,
     size=13, color=CYAN2, bold=True)
text(slide, "A negative overall cannot identify the failure; retrieve all six metrics by context.",
     7.24, 5.53, 4.62, 0.34, size=11.5, color=WHITE)
finish(slide, note)


# 19
note = (
    "Introduce the six metrics as distinct diagnostic lenses. PDS and MSE mainly evaluate group-level effects; the four DE metrics also "
    "depend on cell-level dispersion because significance is estimated from 400 cells."
)
slide = new_slide("Evaluation", "Six metrics = six failure modes",
                  "A single overall score hides whether identity, magnitude, direction, significance, or diversity failed.",
                  "Source: VCC 2026 metric specification", note)
metrics = [
    ("PDS", "Is the perturbation distinguishable?", "rank / direction", CYAN),
    ("MSE", "Is the group expression accurate?", "centroid / magnitude", PURPLE),
    ("FID", "Are significant directions correct?", "sign + coverage", PINK),
    ("REACH", "How deep is the 90%-correct prefix?", "confidence ranking", GREEN),
    ("JAC", "Do significant gene sets overlap?", "DE calibration", YELLOW),
    ("NMAE", "Are reference-significant LFCs accurate?", "effect size", BLUE),
]
for i, (abbr, question, focus, col) in enumerate(metrics):
    x = 0.70 + (i % 3) * 4.10
    y = 1.70 + (i // 3) * 2.18
    rect(slide, x, y, 3.82, 1.78, fill=SURFACE, line=col)
    label(slide, abbr, x + 0.25, y + 0.24, 0.78, fill=col, color=BG if col in [CYAN, GREEN, YELLOW] else WHITE)
    text(slide, question, x + 0.25, y + 0.73, 3.32, 0.52, size=13, bold=True)
    text(slide, focus, x + 0.25, y + 1.34, 3.32, 0.24, size=10, color=MUTED, font=MONO)
rect(slide, 1.06, 6.13, 11.20, 0.45, fill=BG2, line=SURFACE2)
text(slide, "Always inspect 3 contexts × 6 metrics. Overall alone is not a debugging signal.",
     1.30, 6.23, 10.72, 0.24, size=12.2, color=RED, bold=True, align=PP_ALIGN.CENTER)
finish(slide, note)


# 20
note = (
    "PDS compares each predicted perturbation delta with the panel of measured deltas and asks where the correct target ranks by cosine distance. "
    "A generic stress response shared across targets will not discriminate perturbations. All panel targets are excluded from the feature axis for PDS."
)
slide = new_slide("Evaluation", "PDS: make each target response identifiable",
                  "Correct direction is useful only if it is more similar to the matching perturbation than to the other 299.",
                  "Source: VCC 2026 metric specification", note)
circle(slide, 1.26, 2.70, 1.05, fill=CYAN, transparency=8)
text(slide, "Δ̂p", 1.26, 2.70, 1.05, 1.05, size=20, color=BG, bold=True, font=MONO,
     align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
text(slide, "predicted delta", 0.88, 3.90, 1.80, 0.28, size=11, color=MUTED, align=PP_ALIGN.CENTER)
arrow_between(slide, 2.58, 3.23, 3.46, 3.23, color=CYAN)
for i in range(12):
    ang = 2 * math.pi * i / 12
    x = 5.30 + 1.42 * math.cos(ang)
    y = 3.18 + 1.04 * math.sin(ang)
    col = GREEN if i == 2 else SURFACE2
    circle(slide, x, y, 0.35, fill=col, line=GREEN if i == 2 else MUTED2)
    text(slide, "p" if i == 2 else str(i + 1), x, y, 0.35, 0.35, size=8.5,
         color=BG if i == 2 else MUTED, bold=(i == 2), align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
text(slide, "300 measured target deltas", 4.26, 4.75, 2.82, 0.30, size=12, bold=True, align=PP_ALIGN.CENTER)
rect(slide, 8.14, 1.85, 4.20, 3.42, fill=SURFACE, line=PURPLE)
text(slide, "Model behavior rewarded", 8.46, 2.15, 3.56, 0.34, size=15, color=PURPLE2, bold=True)
bullet_list(slide, [
    "target-specific multi-gene direction",
    "correct relative geometry across the panel",
    "robustness after excluding all 300 target genes",
    "avoidance of one generic response for every target"
], 8.46, 2.72, 3.44, 1.86, size=11.8, bullet_color=CYAN)
rect(slide, 1.10, 5.68, 11.05, 0.58, fill=BG2, line=RED)
text(slide, "Failure signature: similar output for every target → reasonable-looking cells, poor discrimination.",
     1.36, 5.83, 10.53, 0.28, size=12.5, color=RED, bold=True, align=PP_ALIGN.CENTER)
finish(slide, note)


# 21
note = (
    "Expression MSE focuses on normalized group-level profiles with sampling-noise correction and a real-effect normalization. "
    "It rewards centroid and magnitude accuracy. It does not by itself guarantee calibrated cell-to-cell variability."
)
slide = new_slide("Evaluation", "Expression MSE: put the centroid in the right place",
                  "A good distribution with the wrong mean still fails the competition.",
                  "Source: VCC 2026 metric specification", note)
rect(slide, 0.80, 1.78, 7.26, 3.90, fill=SURFACE, line=SURFACE2)
text(slide, "Scorer-space transformation", 1.10, 2.06, 2.88, 0.32, size=15, bold=True)
steps = [
    ("sum", "400 cells", CYAN), ("normalize", "to 50,000", PURPLE),
    ("log1p", "group profile", PINK), ("compare", "prediction vs truth", GREEN),
]
for i, (ttl, body, col) in enumerate(steps):
    x = 1.12 + i * 1.65
    circle(slide, x + 0.40, 2.72, 0.76, fill=col, transparency=10)
    text(slide, str(i + 1), x + 0.40, 2.72, 0.76, 0.76, size=16, color=BG if col in [CYAN, GREEN] else WHITE,
         bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    text(slide, ttl, x, 3.66, 1.56, 0.28, size=11, bold=True, align=PP_ALIGN.CENTER)
    text(slide, body, x, 4.02, 1.56, 0.26, size=9.5, color=MUTED, align=PP_ALIGN.CENTER)
    if i < 3:
        arrow_between(slide, x + 1.22, 3.10, x + 1.55, 3.10, color=MUTED2)
text(slide, "sampling-noise corrected • normalized by true perturbation distance", 1.10, 4.80, 6.66, 0.30,
     size=10.5, color=YELLOW2, font=MONO, align=PP_ALIGN.CENTER)
rect(slide, 8.40, 1.78, 3.98, 1.63, fill=SURFACE, line=GREEN)
text(slide, "Protects", 8.70, 2.07, 1.04, 0.30, size=12, color=GREEN, bold=True)
text(slide, "centroid location\neffect magnitude", 9.76, 2.03, 2.12, 0.74, size=15, bold=True)
rect(slide, 8.40, 3.63, 3.98, 2.05, fill=SURFACE, line=RED)
text(slide, "Does not guarantee", 8.70, 3.92, 1.58, 0.30, size=12, color=RED, bold=True)
text(slide, "realistic variance\ncorrect DE call count\nresponder heterogeneity", 8.70, 4.36, 3.12, 0.94, size=13, bold=True)
text(slide, "First optimize Δ̂c,p. Then calibrate the 400-cell distribution around it.",
     1.18, 6.10, 10.98, 0.34, size=14, color=CYAN2, bold=True, align=PP_ALIGN.CENTER)
finish(slide, note)


# 22
note = (
    "The four DE metrics are complementary. Fidelity balances sign correctness with coverage; reach measures a high-confidence prefix; "
    "Jaccard compares significant sets; NMAE measures fold-change magnitude on reference-significant genes."
)
slide = new_slide("Evaluation", "Four DE metrics probe different mistakes",
                  "Sign, confidence ordering, significance calibration, and effect size are not interchangeable.",
                  "Source: VCC 2026 metric specification", note)
simple_table(slide, 0.70, 1.70, [1.38, 2.45, 2.48, 2.45, 2.45], 0.68,
             ["Metric", "Question", "Rewards", "Punishes", "Model lever"],
             [
                 ["FID", "Are DE directions right?", "sign + adequate coverage", "few calls or wrong signs", "direction-aware loss"],
                 ["REACH", "How deep stays ≥90% right?", "calibrated confidence rank", "bad early ranking", "rank / margin loss"],
                 ["JAC", "Do significant sets overlap?", "matched DE cardinality", "wrong variance / FDR", "dispersion + threshold"],
                 ["NMAE", "Are log2FC values right?", "magnitude calibration", "over/under-shoot", "effect scaling"],
             ], font_size=9.6, accents=[PINK, GREEN, YELLOW, BLUE])
rect(slide, 0.98, 5.52, 11.34, 0.78, fill=BG2, line=PURPLE)
text(slide, "Common dependency", 1.28, 5.77, 1.62, 0.28, size=11, color=PURPLE2, bold=True)
text(slide, "DE is called from predicted cells versus real held-out controls using per-gene Wilcoxon tests and BH-FDR.",
     3.02, 5.72, 8.90, 0.37, size=12.5, bold=True)
finish(slide, note)


# 23
note = (
    "Explain the central multi-objective tension. Amplifying deltas can help reach but hurt MSE/NMAE. Shrinking everything protects MSE "
    "but destroys discrimination. Excess variance hides DE; insufficient variance over-calls it. Calibration and ensembling are first-class components."
)
slide = new_slide("Evaluation", "The metrics pull in different directions",
                  "Winning requires calibrated trade-offs, not a single loss optimized harder.",
                  "Interpretation of official metrics", note)
# Triangle
pts = [(3.45, 1.86), (1.20, 5.20), (5.70, 5.20)]
line(slide, *pts[0], *pts[1], color=CYAN, width=2.5)
line(slide, *pts[1], *pts[2], color=PURPLE, width=2.5)
line(slide, *pts[2], *pts[0], color=PINK, width=2.5)
for (x, y), ttl, col in zip(pts, ["Mean accuracy", "Target identity", "DE calibration"], [CYAN, PURPLE, PINK]):
    circle(slide, x - 0.38, y - 0.38, 0.76, fill=col)
    text(slide, ttl, x - 0.88, y + 0.52 if y < 3 else y + 0.48, 1.76, 0.40, size=11.5,
         color=WHITE, bold=True, align=PP_ALIGN.CENTER)
circle(slide, 3.07, 3.39, 0.76, fill=GREEN)
text(slide, "calibrated\nmodel", 2.72, 3.29, 1.46, 0.94, size=11, color=BG, bold=True,
     align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
rect(slide, 7.02, 1.76, 5.36, 3.90, fill=SURFACE, line=SURFACE2)
text(slide, "Typical interventions and side effects", 7.36, 2.06, 4.68, 0.36, size=15, bold=True)
simple_table(slide, 7.34, 2.60, [2.18, 2.56], 0.56,
             ["Intervention", "Possible side effect"],
             [
                 ["Increase guidance", "MSE / NMAE overshoot"],
                 ["Shrink all deltas", "PDS / reach collapse"],
                 ["Reduce dispersion", "too many DE calls"],
                 ["Add diversity", "true DE becomes underpowered"],
             ], font_size=9.3, accents=[CYAN, PURPLE, PINK, GREEN])
text(slide, "Therefore: use metric-aware validation and choose a Pareto-robust ensemble.",
     1.16, 6.13, 11.00, 0.34, size=14, color=YELLOW2, bold=True, align=PP_ALIGN.CENTER)
finish(slide, note)


# 24
note = (
    "This hierarchy prevents premature optimization. File validity is foundational. Target-specific mean effects come next. Direction and magnitude "
    "follow, then cell-level calibration, and only then expensive generative or RL refinements."
)
slide = new_slide("Evaluation", "A scoring-aware hierarchy of needs",
                  "Do not optimize the top layer while the lower layer is still broken.",
                  "Team decision framework", note)
levels = [
    ("1", "VALID OUTPUT", "axes • labels • 400 cells • raw sparse counts", CYAN, 10.6),
    ("2", "TARGET-SPECIFIC MEAN", "correct perturbation identity and centroid", PURPLE, 9.5),
    ("3", "DIRECTION + MAGNITUDE", "DE signs, confidence rank, calibrated LFC", PINK, 8.4),
    ("4", "CELL DISTRIBUTION", "depth, variance, zeros, responder mixtures", GREEN, 7.3),
    ("5", "ADVANCED OPTIMIZATION", "diffusion / flow refinement • optional RL", YELLOW, 6.2),
]
for i, (num, ttl, body, col, width) in enumerate(levels):
    x = (W - width) / 2
    y = 5.83 - i * 0.90
    rect(slide, x, y, width, 0.68, fill=col, transparency=8)
    circle(slide, x + 0.16, y + 0.11, 0.46, fill=BG, line=WHITE)
    text(slide, num, x + 0.16, y + 0.11, 0.46, 0.46, size=11, color=WHITE, bold=True,
         align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    text(slide, ttl, x + 0.78, y + 0.10, 2.82, 0.25, size=11.5, color=BG if col in [CYAN, GREEN, YELLOW] else WHITE, bold=True)
    text(slide, body, x + 3.42, y + 0.10, width - 3.70, 0.28, size=9.7,
         color=BG if col in [CYAN, GREEN, YELLOW] else WHITE, align=PP_ALIGN.RIGHT)
finish(slide, note)


# 25
note = (
    "Before deep models, build a baseline ladder with increasing biological specificity. Each rung isolates value added. A sophisticated model that "
    "cannot beat similarity-weighted low-rank transfer on held-out contexts should not be promoted."
)
slide = new_slide("Model strategy", "Baselines before deep models",
                  "Every rung answers a different question and creates a safe fallback.",
                  "Team recommendation", note)
rungs = [
    ("B0", "Format-only null", "Pipeline passes?", MUTED2, 1.30),
    ("B1", "Control resampling", "No-response reference?", CYAN, 1.90),
    ("B2", "Global mean response", "Public data beats zero?", PURPLE, 2.50),
    ("B3", "Target transfer", "Priors predict target?", PINK, 3.10),
    ("B4", "Context-conditioned low rank", "Context interaction?", GREEN, 3.70),
    ("B5", "Hybrid count generator", "Calibrated variability?", YELLOW, 4.30),
]
for i, (num, ttl, q, col, y) in enumerate(rungs):
    x = 0.90 + i * 0.43
    width = 7.55 - i * 0.43
    rect(slide, x, y, width, 0.48, fill=col, transparency=8)
    text(slide, num, x + 0.14, y + 0.08, 0.56, 0.24, size=9.5, color=BG if col in [CYAN, GREEN, YELLOW] else WHITE, bold=True)
    text(slide, ttl, x + 0.78, y + 0.07, 3.35, 0.26, size=10.7, color=BG if col in [CYAN, GREEN, YELLOW] else WHITE, bold=True)
    text(slide, q, x + 4.10, y + 0.07, width - 4.30, 0.26, size=9.7,
         color=BG if col in [CYAN, GREEN, YELLOW] else WHITE, align=PP_ALIGN.RIGHT)
rect(slide, 8.82, 1.64, 3.54, 3.22, fill=SURFACE, line=PURPLE)
text(slide, "Promotion gate", 9.15, 1.94, 2.88, 0.32, size=15, color=PURPLE2, bold=True, align=PP_ALIGN.CENTER)
text(slide, "Mean held-out gain\n+\nworst-context safety\n+\nreproducible artifact", 9.22, 2.58, 2.74, 1.66,
     size=14, bold=True, align=PP_ALIGN.CENTER)
text(slide, "Do not compare models on training cell lines or random cell splits.",
     1.26, 6.17, 10.80, 0.32, size=13, color=RED, bold=True, align=PP_ALIGN.CENTER)
finish(slide, note)


# 26
note = (
    "Use a decomposed response: context baseline plus a global target effect, a context-target interaction, and stochastic residual. "
    "Low rank and shrinkage make the interaction estimable. External data provide perturbation supervision; challenge controls provide context."
)
slide = new_slide("Model strategy", "A scientific decomposition of the response",
                  "Separate what can transfer from what must adapt to the new control population.",
                  "Team modeling framework", note)
terms = [
    ("x⁰c,i", "control cell", "observed context state", CYAN),
    ("Δglobal(p)", "shared target effect", "learned across public cell lines", PURPLE),
    ("Δint(c,p)", "context interaction", "low-rank, similarity-weighted", PINK),
    ("εc,p,i", "cell residual", "depth + dispersion + state", GREEN),
]
for i, (sym, ttl, body, col) in enumerate(terms):
    x = 0.64 + i * 3.12
    rect(slide, x, 1.92, 2.74, 2.45, fill=SURFACE, line=col)
    text(slide, sym, x + 0.18, 2.24, 2.38, 0.62, size=23, color=col, bold=True, font=MONO, align=PP_ALIGN.CENTER)
    text(slide, ttl, x + 0.18, 3.04, 2.38, 0.32, size=13, bold=True, align=PP_ALIGN.CENTER)
    text(slide, body, x + 0.24, 3.52, 2.26, 0.48, size=10.5, color=MUTED, align=PP_ALIGN.CENTER)
    if i < 3:
        text(slide, "+", x + 2.80, 2.82, 0.30, 0.45, size=22, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
rect(slide, 1.18, 4.96, 10.98, 0.90, fill=BG2, line=PURPLE)
text(slide, "x̂¹c,p,i  =  x⁰c,i  ⊕  Δglobal(p)  ⊕  Δint(c,p)  ⊕  εc,p,i",
     1.44, 5.22, 10.46, 0.36, size=18, color=PURPLE2, bold=True, font=MONO, align=PP_ALIGN.CENTER)
text(slide, "⊕ denotes a count-safe transformation—not unrestricted addition in raw count space.",
     2.08, 6.16, 9.18, 0.28, size=10.8, color=MUTED, align=PP_ALIGN.CENTER)
finish(slide, note)


# 27
note = (
    "The challenge controls contain no matched perturbations, so supervised signal must come from licensed external data. Build a provenance registry, "
    "harmonize gene symbols and modality, and measure transfer under held-out cell-line splits. Source diversity matters more than raw cell count alone."
)
slide = new_slide("Model strategy", "External data are the training set",
                  "Use perturbation outcomes for supervision; use challenge controls only for zero-shot context conditioning.",
                  "Sources: VCC 2026 Data; VCC Rules; public Perturb-seq literature", note)
sources = [
    ("2025 H1 hESC", "same challenge lineage; useful reference", CYAN),
    ("Public CRISPRi Perturb-seq", "single-gene response supervision", PURPLE),
    ("Genome-scale screens", "target coverage and response bases", PINK),
    ("Atlas / bulk signatures", "context and pathway priors", GREEN),
]
for i, (ttl, body, col) in enumerate(sources):
    y = 1.70 + i * 1.06
    circle(slide, 0.82, y + 0.12, 0.56, fill=col)
    text(slide, str(i + 1), 0.82, y + 0.12, 0.56, 0.56, size=12, color=BG if col in [CYAN, GREEN] else WHITE,
         bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    text(slide, ttl, 1.62, y + 0.08, 2.90, 0.30, size=13, bold=True)
    text(slide, body, 4.44, y + 0.08, 2.88, 0.36, size=10.5, color=MUTED)
arrow_between(slide, 7.50, 3.35, 8.18, 3.35, color=CYAN)
rect(slide, 8.30, 1.70, 4.14, 4.10, fill=SURFACE, line=CYAN)
text(slide, "Harmonization gates", 8.62, 2.00, 3.50, 0.34, size=15, color=CYAN2, bold=True)
bullet_list(slide, [
    "gene symbol and 18,533-axis mapping",
    "CRISPRi versus knockout / activation",
    "time point and knockdown strength",
    "library depth and count technology",
    "cell-line similarity and source ablation",
    "license and provenance record"
], 8.62, 2.52, 3.40, 2.56, size=11.2, bullet_color=PURPLE2, gap=3)
rect(slide, 1.02, 6.12, 11.28, 0.42, fill=BG2, line=RED)
text(slide, "Never paste measured final outcomes; all final predictions must be model-generated under the competition rules.",
     1.25, 6.21, 10.82, 0.22, size=11.5, color=RED, bold=True, align=PP_ALIGN.CENTER)
finish(slide, note)


# 28
note = (
    "This is the recommended system. A context encoder summarizes NTC cells. A perturbation encoder represents unseen genes using multi-view priors. "
    "A mean-effect predictor produces a calibrated delta. A count generator samples 400 cells anchored to real controls. The local scorer calibrates and blends candidates."
)
slide = new_slide("Model strategy", "Recommended hybrid system",
                  "Make the deterministic mean effect the spine; make generative modeling a calibrated distribution layer.",
                  "Team recommendation", note)
blocks = [
    (0.55, 1.88, 2.00, 1.45, "NTC controls", "context distribution", CYAN),
    (0.55, 4.18, 2.00, 1.45, "Target gene", "sequence • graph • text", PURPLE),
    (3.03, 1.88, 2.05, 1.45, "Context encoder", "set / pseudobulk / modules", CYAN),
    (3.03, 4.18, 2.05, 1.45, "Perturbation encoder", "multi-view gene embedding", PURPLE),
    (5.56, 2.76, 2.28, 1.75, "Mean-effect predictor", "Δglobal + Δinteraction", PINK),
    (8.34, 2.76, 2.02, 1.75, "Count generator", "NB/DM or residual flow", GREEN),
    (10.84, 2.76, 1.94, 1.75, "400 cells", "raw sparse counts", YELLOW),
]
for x, y, w, h, ttl, body, col in blocks:
    rect(slide, x, y, w, h, fill=SURFACE, line=col)
    text(slide, ttl, x + 0.14, y + 0.28, w - 0.28, 0.42, size=13.2, bold=True, align=PP_ALIGN.CENTER)
    text(slide, body, x + 0.14, y + 0.84, w - 0.28, 0.38, size=9.5, color=MUTED, align=PP_ALIGN.CENTER)
arrow_between(slide, 2.66, 2.61, 2.94, 2.61, color=CYAN)
arrow_between(slide, 2.66, 4.91, 2.94, 4.91, color=PURPLE)
arrow_between(slide, 5.18, 2.61, 5.48, 3.25, color=CYAN)
arrow_between(slide, 5.18, 4.91, 5.48, 4.03, color=PURPLE)
arrow_between(slide, 7.93, 3.63, 8.23, 3.63, color=PINK)
arrow_between(slide, 10.46, 3.63, 10.75, 3.63, color=GREEN)
rect(slide, 4.12, 5.94, 5.14, 0.52, fill=BG2, line=CYAN)
text(slide, "exact metric emulator  →  calibration  →  ensemble", 4.34, 6.07, 4.70, 0.25,
     size=11.5, color=CYAN2, bold=True, font=MONO, align=PP_ALIGN.CENTER)
line(slide, 6.68, 4.66, 6.68, 5.84, color=MUTED2, width=1.5, dash=True)
finish(slide, note)


# 29
note = (
    "The context encoder should accept a set or distribution of cells and be insensitive to their order. Start with robust pseudobulk and module summaries; "
    "then compare learned set encoders. Include nuisance statistics explicitly so biology is not confounded with depth or sparsity."
)
slide = new_slide("Model components", "Context encoder: summarize a population",
                  "A robust distributional representation is safer than a context-ID embedding.",
                  "Team recommendation", note)
features = [
    ("Basal state", "log-normalized pseudobulk", CYAN),
    ("Cell states", "mixture proportions / latent clusters", PURPLE),
    ("Gene programs", "pathway and module activities", PINK),
    ("Covariance", "low-rank co-expression structure", GREEN),
    ("Technical state", "depth, detected genes, zero fraction", YELLOW),
    ("Uncertainty", "bootstrap variability across NTC cells", BLUE),
]
for i, (ttl, body, col) in enumerate(features):
    x = 0.68 + (i % 3) * 4.12
    y = 1.72 + (i // 3) * 1.66
    rect(slide, x, y, 3.78, 1.34, fill=SURFACE, line=col)
    circle(slide, x + 0.20, y + 0.23, 0.48, fill=col)
    text(slide, str(i + 1), x + 0.20, y + 0.23, 0.48, 0.48, size=11, color=BG if col in [CYAN, GREEN, YELLOW] else WHITE,
         bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    text(slide, ttl, x + 0.84, y + 0.22, 2.62, 0.31, size=13, bold=True)
    text(slide, body, x + 0.84, y + 0.65, 2.62, 0.40, size=10.3, color=MUTED)
rect(slide, 1.02, 5.32, 11.26, 0.92, fill=BG2, line=PURPLE)
text(slide, "Recommended ablation order", 1.30, 5.54, 2.16, 0.28, size=11, color=PURPLE2, bold=True)
text(slide, "pseudobulk only  →  + modules  →  + state mixture  →  + learned set encoder",
     3.68, 5.48, 8.18, 0.36, size=13.2, color=WHITE, bold=True, font=MONO)
text(slide, "Promotion criterion: gain on unseen cell lines, not reconstruction of the training context label.",
     3.68, 5.92, 8.18, 0.24, size=10.2, color=MUTED)
finish(slide, note)


# 30
note = (
    "The final target panel changes, so the perturbation encoder must generalize by biology. Combine learned perturbation signatures with priors such as "
    "protein/gene embeddings, pathways, regulatory graphs, and basal target expression. Missing priors should reduce confidence rather than produce arbitrary effects."
)
slide = new_slide("Model components", "Perturbation encoder: represent an unseen gene",
                  "Target IDs should be endpoints into biological structure—not free lookup tokens.",
                  "Sources: GEARS; CPA; team recommendation", note)
views = [
    ("Functional text", "gene / protein embedding", PURPLE),
    ("Pathways", "GO and curated gene sets", CYAN),
    ("Regulatory graph", "TF, PPI, co-expression edges", PINK),
    ("Sequence / protein", "domain and family similarity", GREEN),
    ("Perturbation data", "learned response signature", YELLOW),
    ("Target context", "basal expression + accessibility proxy", BLUE),
]
for i, (ttl, body, col) in enumerate(views):
    ang = 2 * math.pi * i / 6 - math.pi / 2
    x = 3.30 + 2.15 * math.cos(ang)
    y = 3.58 + 1.75 * math.sin(ang)
    rect(slide, x - 0.88, y - 0.40, 1.76, 0.80, fill=SURFACE, line=col)
    text(slide, ttl, x - 0.76, y - 0.25, 1.52, 0.24, size=9.8, bold=True, align=PP_ALIGN.CENTER)
    text(slide, body, x - 0.76, y + 0.03, 1.52, 0.24, size=7.8, color=MUTED, align=PP_ALIGN.CENTER)
    line(slide, x, y, 3.30, 3.58, color=MUTED2, width=1.0)
circle(slide, 2.57, 2.85, 1.46, fill=PURPLE, transparency=8)
text(slide, "zp", 2.57, 2.85, 1.46, 1.46, size=28, bold=True, font=MONO,
     align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
rect(slide, 7.05, 1.78, 5.32, 3.90, fill=SURFACE, line=PURPLE)
text(slide, "Fusion rules", 7.38, 2.08, 4.66, 0.32, size=15, color=PURPLE2, bold=True)
bullet_list(slide, [
    "freeze or regularize general-purpose embeddings",
    "learn small modality-specific adapters",
    "mask target identity during some training batches",
    "evaluate targets held out by gene family / pathway",
    "emit uncertainty when priors disagree or coverage is low"
], 7.38, 2.60, 4.48, 2.34, size=11.5, bullet_color=CYAN)
text(slide, "Key test: entirely unseen target genes in entirely unseen cell lines.",
     7.38, 5.14, 4.48, 0.28, size=10.8, color=YELLOW2, bold=True)
finish(slide, note)


# 31
note = (
    "Predict the perturbation effect in a low-dimensional response basis, then decode to genes. Train with pseudobulk and direction-aware losses. "
    "Use shrinkage and calibrated interaction terms to avoid overfitting sparse context-target evidence."
)
slide = new_slide("Model components", "Mean-effect predictor: optimize the signal first",
                  "A low-rank response basis reduces 18,533 outputs to a transferable biological coordinate system.",
                  "Team recommendation; related ideas: CPA and GEARS", note)
rect(slide, 0.72, 1.82, 2.26, 1.44, fill=SURFACE, line=CYAN)
text(slide, "zc", 1.18, 2.17, 1.34, 0.48, size=24, color=CYAN2, bold=True, font=MONO, align=PP_ALIGN.CENTER)
text(slide, "context embedding", 0.98, 2.76, 1.74, 0.24, size=9.5, color=MUTED, align=PP_ALIGN.CENTER)
rect(slide, 0.72, 4.16, 2.26, 1.44, fill=SURFACE, line=PURPLE)
text(slide, "zp", 1.18, 4.51, 1.34, 0.48, size=24, color=PURPLE2, bold=True, font=MONO, align=PP_ALIGN.CENTER)
text(slide, "target embedding", 0.98, 5.10, 1.74, 0.24, size=9.5, color=MUTED, align=PP_ALIGN.CENTER)
arrow_between(slide, 3.10, 2.56, 4.10, 3.42, color=CYAN)
arrow_between(slide, 3.10, 4.88, 4.10, 4.02, color=PURPLE)
rect(slide, 4.20, 2.72, 2.38, 1.92, fill=BG2, line=PINK)
text(slide, "interaction\nnetwork", 4.50, 3.17, 1.78, 0.76, size=18, color=PINK2, bold=True, align=PP_ALIGN.CENTER)
arrow_between(slide, 6.70, 3.68, 7.42, 3.68, color=PINK)
rect(slide, 7.54, 2.72, 2.16, 1.92, fill=SURFACE, line=GREEN)
text(slide, "K-dimensional\nresponse basis", 7.82, 3.16, 1.60, 0.78, size=16, color=GREEN2, bold=True, align=PP_ALIGN.CENTER)
arrow_between(slide, 9.82, 3.68, 10.46, 3.68, color=GREEN)
rect(slide, 10.56, 2.72, 2.04, 1.92, fill=SURFACE, line=YELLOW)
text(slide, "Δ̂c,p\n18,533 genes", 10.80, 3.15, 1.56, 0.80, size=16, color=YELLOW2, bold=True, align=PP_ALIGN.CENTER)
rect(slide, 2.28, 5.94, 8.76, 0.48, fill=BG2, line=SURFACE2)
text(slide, "Loss = centroid + direction/rank + LFC calibration + low-rank/shrinkage regularization",
     2.52, 6.06, 8.28, 0.24, size=11.3, color=CYAN2, bold=True, font=MONO, align=PP_ALIGN.CENTER)
finish(slide, note)


# 32
note = (
    "Convert the mean profile into 400 raw-count cells. A pragmatic first choice is control-anchored negative-binomial or Dirichlet-multinomial sampling: "
    "sample a real control state, apply a positive mean transformation, draw a library size, then draw counts with calibrated dispersion."
)
slide = new_slide("Model components", "Generate raw counts without losing the mean",
                  "A count-safe reconstruction layer bridges normalized modeling space and the submission contract.",
                  "Team recommendation", note)
steps = [
    ("Sample", "a real NTC cell or latent state", CYAN),
    ("Transform", "apply positive target-specific mean shift", PURPLE),
    ("Calibrate", "library size, dispersion, zero rate", PINK),
    ("Draw", "NB / DM counts or residual-flow sample", GREEN),
    ("Validate", "integer • sparse • 400 cells", YELLOW),
]
for i, (ttl, body, col) in enumerate(steps):
    x = 0.58 + i * 2.53
    circle(slide, x + 0.74, 1.82, 0.72, fill=col)
    text(slide, str(i + 1), x + 0.74, 1.82, 0.72, 0.72, size=15, color=BG if col in [CYAN, GREEN, YELLOW] else WHITE,
         bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    text(slide, ttl, x, 2.78, 2.20, 0.34, size=14, bold=True, align=PP_ALIGN.CENTER)
    text(slide, body, x, 3.24, 2.20, 0.64, size=10.5, color=MUTED, align=PP_ALIGN.CENTER)
    if i < 4:
        arrow_between(slide, x + 1.92, 2.18, x + 2.34, 2.18, color=MUTED2)
rect(slide, 0.96, 4.52, 11.42, 1.08, fill=SURFACE, line=PURPLE)
text(slide, "Recommended first implementation", 1.24, 4.78, 2.80, 0.30, size=12, color=PURPLE2, bold=True)
text(slide, "control-anchored mean transform  +  gene-wise NB dispersion  +  empirical library-size sampling",
     4.10, 4.70, 7.88, 0.42, size=13.2, bold=True, font=MONO)
text(slide, "Upgrade to residual flow/diffusion only if held-out DE calibration improves.",
     4.10, 5.15, 7.88, 0.26, size=10.4, color=MUTED)
rect(slide, 2.08, 6.00, 9.16, 0.42, fill=BG2, line=RED)
text(slide, "Never round early: preserve expected means, then draw or round only at the final count layer.",
     2.32, 6.10, 8.68, 0.22, size=11.2, color=RED, bold=True, align=PP_ALIGN.CENTER)
finish(slide, note)


# 33
note = (
    "Diffusion or flow matching is useful for the residual distribution after a reliable mean effect exists. Work in a compact latent or residual space, "
    "condition on context and target, and reconstruct counts through a likelihood-aware decoder. Compare against a simpler NB/DM sampler."
)
slide = new_slide("Model choices", "Where diffusion or flow matching helps",
                  "Use it to model residual heterogeneity—not to discover the mean effect from hidden challenge labels.",
                  "Team recommendation", note)
rect(slide, 0.72, 1.76, 7.18, 3.98, fill=SURFACE, line=PURPLE)
text(slide, "Conditional residual generator", 1.05, 2.04, 3.76, 0.34, size=16, color=PURPLE2, bold=True)
for i in range(6):
    x = 1.18 + i * 0.78
    d = 0.60 - i * 0.06
    circle(slide, x, 2.89 + (i % 2) * 0.18, d, fill=[MUTED2, BLUE, PURPLE, PINK, CYAN, GREEN][i], transparency=12)
    if i < 5:
        arrow_between(slide, x + d + 0.05, 3.20, x + 0.70, 3.20, color=MUTED2)
text(slide, "noise / residual", 1.10, 3.84, 1.62, 0.28, size=10, color=MUTED)
text(slide, "context + target guided trajectory", 3.13, 3.84, 2.72, 0.28, size=10, color=MUTED, align=PP_ALIGN.CENTER)
text(slide, "cell residual", 6.10, 3.84, 1.34, 0.28, size=10, color=MUTED, align=PP_ALIGN.RIGHT)
text(slide, "εθ(z, t | zc, zp, Δ̂c,p)", 1.64, 4.58, 5.40, 0.42, size=19, color=CYAN2, bold=True, font=MONO, align=PP_ALIGN.CENTER)
rect(slide, 8.26, 1.76, 4.14, 1.72, fill=SURFACE, line=GREEN)
text(slide, "Use when", 8.56, 2.04, 1.34, 0.30, size=12, color=GREEN, bold=True)
bullet_list(slide, ["NB/DM underfits covariance", "held-out Jaccard improves", "diversity remains calibrated"],
            8.56, 2.48, 3.34, 0.78, size=10.8, bullet_color=GREEN, gap=2)
rect(slide, 8.26, 3.76, 4.14, 1.98, fill=SURFACE, line=RED)
text(slide, "Stop when", 8.56, 4.04, 1.34, 0.30, size=12, color=RED, bold=True)
bullet_list(slide, ["centroid drifts", "mode diversity hurts DE power", "no LOCO gain versus count sampler"],
            8.56, 4.48, 3.34, 0.86, size=10.8, bullet_color=RED, gap=2)
text(slide, "Diffusion is a distributional upgrade, not a prerequisite.",
     1.24, 6.16, 10.84, 0.32, size=14, color=YELLOW2, bold=True, align=PP_ALIGN.CENTER)
finish(slide, note)


# 34
note = (
    "RL is not the core method because there is no sequential environment and no challenge ground truth for A–F. The action space is enormous and the reward is "
    "a small set of panel-coupled metrics. If used, RL should update a small adapter against exact offline metrics on public held-out contexts with KL and diversity constraints."
)
slide = new_slide("Model choices", "Where reinforcement learning helps—and why it is optional",
                  "The problem is supervised conditional distribution estimation, not sequential control.",
                  "Team recommendation", note)
rect(slide, 0.72, 1.72, 5.56, 4.08, fill=SURFACE, line=RED)
label(slide, "Why not core", 1.02, 2.00, 1.25, fill=RED)
bullet_list(slide, [
    "The target is given; the model chooses no intervention.",
    "There is no natural action–state trajectory.",
    "A–F perturbation truth is unavailable during training.",
    "Millions of count outputs receive six scalar rewards.",
    "PDS and DE metrics are panel-coupled and partly discontinuous.",
    "Imperfect verifiers invite reward hacking and mode collapse."
], 1.04, 2.56, 4.74, 2.72, size=11.4, bullet_color=RED, gap=3)
rect(slide, 6.66, 1.72, 5.70, 4.08, fill=SURFACE, line=GREEN)
label(slide, "Optional lane", 6.96, 2.00, 1.20, fill=GREEN, color=BG)
bullet_list(slide, [
    "Freeze a supervised generator first.",
    "Train a surrogate/exact metric verifier on public folds.",
    "Update only a small adapter or guidance policy.",
    "Constrain KL, count moments, and diversity.",
    "Evaluate on untouched cell lines and three seeds.",
    "Terminate if any metric or worst context collapses."
], 6.98, 2.56, 4.84, 2.72, size=11.4, bullet_color=GREEN, gap=3)
rect(slide, 1.52, 6.14, 10.28, 0.42, fill=BG2, line=YELLOW)
text(slide, "Compute budget cap: ≤10% for RL until the supervised entry gate is passed.",
     1.80, 6.24, 9.72, 0.22, size=11.5, color=YELLOW2, bold=True, align=PP_ALIGN.CENTER)
finish(slide, note)


# 35
note = (
    "Random cell splits leak cell-line identity and greatly overstate transfer. Use leave-one-cell-line-out outer folds and target-held-out inner folds. "
    "The exact metric emulator must be applied to generated 400-cell groups, not only model-space losses."
)
slide = new_slide("Validation", "LOCO validation is the scientific backbone",
                  "Hold out entire cell contexts; never validate by random cells from a known line.",
                  "Team recommendation", note)
folds = [
    ("Fold 1", "Train: B C D …", "Test: A", CYAN),
    ("Fold 2", "Train: A C D …", "Test: B", PURPLE),
    ("Fold 3", "Train: A B D …", "Test: C", PINK),
]
for i, (ttl, train, test, col) in enumerate(folds):
    x = 0.80 + i * 4.10
    rect(slide, x, 1.84, 3.72, 2.10, fill=SURFACE, line=col)
    label(slide, ttl, x + 0.24, 2.08, 0.82, fill=col, color=BG if col == CYAN else WHITE)
    text(slide, train, x + 0.26, 2.70, 3.20, 0.34, size=12.5, bold=True, font=MONO)
    text(slide, test, x + 0.26, 3.20, 3.20, 0.34, size=13.5, color=col, bold=True, font=MONO)
rect(slide, 0.80, 4.35, 11.92, 1.34, fill=BG2, line=SURFACE2)
criteria = [
    ("Outer split", "unseen cell line", CYAN), ("Inner split", "unseen targets / families", PURPLE),
    ("Scoring", "exact six-metric emulator", PINK), ("Report", "mean + worst context + seeds", GREEN),
]
for i, (ttl, body, col) in enumerate(criteria):
    x = 1.05 + i * 2.93
    text(slide, ttl, x, 4.66, 2.50, 0.28, size=10.5, color=col, bold=True, align=PP_ALIGN.CENTER)
    text(slide, body, x, 5.08, 2.50, 0.30, size=11.3, bold=True, align=PP_ALIGN.CENTER)
text(slide, "Leaderboard feedback is a sparse external check—not the training reward.",
     1.60, 6.16, 10.10, 0.32, size=13.5, color=RED, bold=True, align=PP_ALIGN.CENTER)
finish(slide, note)


# 36
note = (
    "Pre-register ablations so each experiment has a hypothesis, target metrics, and stop rule. This table is a suggested initial sequence. "
    "A component is promoted only if it improves multiple held-out contexts without catastrophic metric regressions."
)
slide = new_slide("Validation", "Ablations and decision gates",
                  "Change one mechanism at a time; attach every run to an expected metric effect.",
                  "Team experiment policy", note)
simple_table(slide, 0.68, 1.68, [1.32, 2.45, 2.55, 2.70, 2.55], 0.60,
             ["ID", "Change", "Primary hypothesis", "Expected metric movement", "Stop rule"],
             [
                 ["A1", "+ target priors", "unseen target geometry", "PDS ↑ • reach ↑", "no LOCO gain"],
                 ["A2", "+ context interaction", "cell-line modulation", "worst context ↑", "variance increases"],
                 ["A3", "+ effect calibration", "correct magnitude", "MSE ↑ • NMAE ↑", "PDS collapses"],
                 ["A4", "+ NB dispersion", "DE call calibration", "JAC ↑ • FID stable", "centroid drifts"],
                 ["A5", "+ residual flow", "covariance / mixtures", "JAC or reach ↑", "no count-model gain"],
                 ["A6", "+ ensemble", "error diversification", "mean + worst ↑", "one fold declines badly"],
                 ["A7", "+ RL adapter", "non-diff metric tuning", "multi-metric ↑", "reward hacking"],
             ], font_size=8.8, accents=[CYAN, PURPLE, PINK, GREEN, YELLOW, BLUE, RED])
text(slide, "Promotion = mean gain + worst-context safety + three-seed stability + valid 360k artifact.",
     1.02, 6.52, 11.30, 0.28, size=12.3, color=CYAN2, bold=True, align=PP_ALIGN.CENTER)
finish(slide, note)


# 37
note = (
    "Assign three durable ownership lanes. P1 owns biological transfer and scientific model selection. P2 owns generative architecture and count calibration. "
    "P3 owns scorer parity, cluster operations, artifacts, and submissions. Every critical workflow has a backup."
)
slide = new_slide("Three-person team", "Three owners, one integrated system",
                  "Clear decision rights reduce duplicated work and prevent silent pipeline failures.",
                  "Team operating recommendation", note)
roles = [
    ("P1", "Biological Modeling Lead", ["public data + provenance", "statistical / low-rank baselines", "context and target priors", "scientific model selection"], CYAN),
    ("P2", "Generative & Optimization Lead", ["flow / diffusion architecture", "count reconstruction + dispersion", "calibration + ensemble", "optional RL lane"], PURPLE),
    ("P3", "Evaluation, MLOps & Competition", ["official metric parity + LOCO", "Slurm / H100 / storage", "artifact registry + reproducibility", "prep, packaging, submission"], GREEN),
]
for i, (pid, ttl, items, col) in enumerate(roles):
    x = 0.68 + i * 4.12
    rect(slide, x, 1.72, 3.78, 3.98, fill=SURFACE, line=col)
    circle(slide, x + 1.37, 1.98, 1.04, fill=col)
    text(slide, pid, x + 1.37, 1.98, 1.04, 1.04, size=19, color=BG if col in [CYAN, GREEN] else WHITE,
         bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    text(slide, ttl, x + 0.32, 3.22, 3.14, 0.56, size=14, bold=True, align=PP_ALIGN.CENTER)
    bullet_list(slide, items, x + 0.36, 4.02, 3.06, 1.32, size=10.7, bullet_color=col, gap=2)
rect(slide, 1.28, 6.05, 10.78, 0.44, fill=BG2, line=RED)
text(slide, "P3 may veto invalid or unreproducible submissions; final model selection requires all three members.",
     1.54, 6.15, 10.26, 0.24, size=11.5, color=RED, bold=True, align=PP_ALIGN.CENTER)
finish(slide, note)


# 38
note = (
    "Use a small RACI table for the critical workstreams. The purpose is not bureaucracy: it ensures that public-data rights, metric parity, "
    "model architecture, and final upload each have exactly one accountable owner."
)
slide = new_slide("Three-person team", "Decision ownership: compact RACI",
                  "One accountable owner per critical workstream; shared review at integration points.",
                  "Team operating recommendation", note)
simple_table(slide, 0.74, 1.68, [4.10, 2.28, 2.28, 2.28], 0.54,
             ["Workstream", "P1 Biology", "P2 Generative", "P3 Eval / MLOps"],
             [
                 ["Public data + licensing", "A / R", "I", "C"],
                 ["Raw-count harmonization", "A / R", "C", "R"],
                 ["LOCO splits + exact scorer", "C", "C", "A / R"],
                 ["Statistical transfer baseline", "A / R", "C", "C"],
                 ["Flow / diffusion generator", "C", "A / R", "C"],
                 ["Count / dispersion calibration", "C", "A / R", "C"],
                 ["Ensemble + model selection", "A", "R", "C"],
                 ["Artifacts + H100 operations", "I", "C", "A / R"],
                 ["Leaderboard / final upload", "A", "C", "R"],
             ], font_size=9.0, accents=[CYAN, PURPLE, PINK, GREEN, YELLOW, BLUE, CYAN, PURPLE, PINK])
text(slide, "R = Responsible   •   A = Accountable   •   C = Consulted   •   I = Informed",
     2.02, 6.79, 9.30, 0.22, size=9.8, color=MUTED, bold=True, align=PP_ALIGN.CENTER)
finish(slide, note)


# 39
note = (
    "The calendar is built backward from October 22 and November 5. The key management choice is to freeze architecture before final controls arrive. "
    "Final phase should be adaptation, calibration, generation, and QC—not research redesign."
)
slide = new_slide("Three-person team", "Timeline to the final submission",
                  "Build evidence in September; freeze in October; preserve recovery time in November.",
                  "Challenge dates from portal snapshot and official documentation", note)
milestones = [
    ("25 Aug", "Audit + recover\ncurrent score", RED),
    ("11 Sep", "Trusted statistical\nbaseline", CYAN),
    ("25 Sep", "Generator\ngo / no-go", PURPLE),
    ("09 Oct", "Ensemble rule\nfrozen", PINK),
    ("21 Oct", "Architecture + code\nfreeze", GREEN),
    ("22 Oct", "D / E / F controls\nreleased", YELLOW),
    ("30 Oct", "Final candidates\ngenerated", BLUE),
    ("04 Nov", "Submit with\nrecovery margin", GREEN),
    ("05 Nov", "23:59 UTC\ndeadline", RED),
]
line(slide, 0.86, 3.35, 12.46, 3.35, color=MUTED2, width=3)
for i, (date, body, col) in enumerate(milestones):
    x = 0.78 + i * 1.40
    above = i % 2 == 0
    circle(slide, x, 3.06, 0.58, fill=col)
    text(slide, str(i + 1), x, 3.06, 0.58, 0.58, size=10, color=BG if col in [CYAN, GREEN, YELLOW] else WHITE,
         bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    yy = 1.72 if above else 4.02
    text(slide, date, x - 0.35, yy, 1.28, 0.28, size=10.5, color=col, bold=True, align=PP_ALIGN.CENTER)
    text(slide, body, x - 0.48, yy + 0.40, 1.54, 0.68, size=9.0, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
rect(slide, 1.30, 5.74, 10.72, 0.72, fill=BG2, line=YELLOW)
text(slide, "Final-phase rule", 1.58, 5.96, 1.54, 0.28, size=10.5, color=YELLOW, bold=True)
text(slide, "No architecture change after release; only context adaptation, calibration, inference, validation, and packaging.",
     3.22, 5.91, 8.38, 0.36, size=11.8, bold=True)
finish(slide, note)


# 40
note = (
    "This is the concrete first two-week plan. It starts by retrieving the current metric breakdown and reproducing the existing artifact, "
    "then locks data provenance and LOCO evaluation, and ends with a trusted baseline plus a single reproducible full-generation rehearsal."
)
slide = new_slide("Three-person team", "The first 14 days",
                  "Recover evidence, establish evaluation, then earn the right to train a large model.",
                  "Team execution plan", note)
weeks = [
    ("Days 1–3", "Forensics", ["retrieve 18 score cells", "reproduce current artifact", "audit A/B/C mapping", "archive checksums"], RED),
    ("Days 4–7", "Foundation", ["public-data inventory", "license + provenance registry", "exact scorer harness", "first LOCO split"], CYAN),
    ("Days 8–11", "Baselines", ["control + generic response", "low-rank target transfer", "context similarity weighting", "metric diagnostics"], PURPLE),
    ("Days 12–14", "Integration", ["count sampler v0", "full 360k rehearsal", "one-command pipeline", "go/no-go review"], GREEN),
]
for i, (days, ttl, items, col) in enumerate(weeks):
    x = 0.60 + i * 3.16
    rect(slide, x, 1.72, 2.86, 4.18, fill=SURFACE, line=col)
    label(slide, days, x + 0.24, 1.98, 1.14, fill=col, color=BG if col in [CYAN, GREEN] else WHITE)
    text(slide, ttl, x + 0.24, 2.50, 2.38, 0.38, size=16, bold=True)
    bullet_list(slide, items, x + 0.26, 3.12, 2.34, 1.78, size=10.8, bullet_color=col, gap=4)
    if i < 3:
        arrow_between(slide, x + 2.92, 3.76, x + 3.10, 3.76, color=MUTED2)
rect(slide, 1.34, 6.20, 10.64, 0.36, fill=BG2, line=YELLOW)
text(slide, "Exit criterion: a non-neural baseline that is trusted scientifically, scored locally, and package-valid end to end.",
     1.56, 6.27, 10.20, 0.20, size=10.8, color=YELLOW2, bold=True, align=PP_ALIGN.CENTER)
finish(slide, note)


# 41
note = (
    "Treat every leaderboard submission as an expensive scientific experiment. Keep immutable artifacts and SHA-256 hashes, use paired comparisons, "
    "record the panel and anchor stamp, and retain the prior champion. Two submissions per day is a ceiling, not a target."
)
slide = new_slide("Competition operations", "Submission protocol: learn without overfitting",
                  "Use the leaderboard for sparse, high-information tests—not iterative reward optimization.",
                  "Source: VCC CLI Guide; team policy", note)
flow = [
    ("Hypothesis", "pre-register expected metric movement", CYAN),
    ("Local proof", "LOCO + exact scorer + QC", PURPLE),
    ("Immutable artifact", "config • commit • SHA-256", PINK),
    ("Two-person review", "P1 + P3 approval", GREEN),
    ("Submit + archive", "six metrics • context • stamp", YELLOW),
]
for i, (ttl, body, col) in enumerate(flow):
    x = 0.60 + i * 2.53
    rect(slide, x, 1.82, 2.18, 2.22, fill=SURFACE, line=col)
    circle(slide, x + 0.73, 2.08, 0.72, fill=col)
    text(slide, str(i + 1), x + 0.73, 2.08, 0.72, 0.72, size=15, color=BG if col in [CYAN, GREEN, YELLOW] else WHITE,
         bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    text(slide, ttl, x + 0.18, 3.00, 1.82, 0.32, size=12, bold=True, align=PP_ALIGN.CENTER)
    text(slide, body, x + 0.18, 3.42, 1.82, 0.42, size=9.2, color=MUTED, align=PP_ALIGN.CENTER)
    if i < 4:
        arrow_between(slide, x + 2.21, 2.96, x + 2.43, 2.96, color=MUTED2)
rect(slide, 0.82, 4.58, 5.62, 1.18, fill=BG2, line=RED)
text(slide, "Do not submit", 1.12, 4.87, 1.42, 0.30, size=11.5, color=RED, bold=True)
text(slide, "identical artifacts • unsupported changes • context-label guesses • DO_NOT_SUBMIT smoke output",
     2.66, 4.80, 3.34, 0.54, size=10.4, bold=True)
rect(slide, 6.82, 4.58, 5.62, 1.18, fill=BG2, line=GREEN)
text(slide, "Retain", 7.12, 4.87, 0.78, 0.30, size=11.5, color=GREEN, bold=True)
text(slide, "prior champion • rollback package • raw JSON scores • model/data/environment provenance",
     8.08, 4.80, 3.92, 0.54, size=10.4, bold=True)
text(slide, "Portal limit: two scored submissions per day, renewing at midnight UTC; one in flight at a time.",
     1.22, 6.20, 10.88, 0.30, size=11.8, color=YELLOW2, bold=True, align=PP_ALIGN.CENTER)
finish(slide, note)


# 42
note = (
    "Review these risks weekly. The most dangerous failures are not always modeling failures: a context swap, scorer mismatch, invalid raw counts, "
    "or a late upload can erase months of scientific work."
)
slide = new_slide("Competition operations", "Risk register: protect the result",
                  "High-impact operational failures deserve the same attention as model error.",
                  "Team risk review", note)
simple_table(slide, 0.72, 1.67, [2.58, 1.42, 2.70, 4.10, 1.18], 0.56,
             ["Risk", "Impact", "Early signal", "Mitigation", "Owner"],
             [
                 ["Context label swap", "Critical", "all metrics near chance", "preserve labels; two-person audit", "P3"],
                 ["Scorer mismatch", "Critical", "local / portal divergence", "pin rule digest; parity tests", "P3"],
                 ["Leaderboard overfit", "High", "portal ↑, LOCO flat", "submission policy; sealed folds", "P1"],
                 ["Wrong DE dispersion", "High", "JAC collapses", "match depth/variance; call-count QC", "P2"],
                 ["Public-data mismatch", "High", "source gains do not transfer", "source ablations; similarity weighting", "P1"],
                 ["Mode collapse", "High", "low within-group variance", "diversity + count-moment checks", "P2"],
                 ["Sparse / storage overflow", "High", "NNZ or quota spike", "CSR; threshold; quota monitor", "P3"],
                 ["Final upload failure", "Critical", "slow / rejected package", "submit by 4 Nov; fallback artifact", "P3"],
             ], font_size=8.7, accents=[RED, RED, YELLOW, PINK, PURPLE, CYAN, GREEN, RED])
text(slide, "Weekly rule: every red risk needs an owner, an observable trigger, and a rehearsed fallback.",
     1.12, 6.60, 11.06, 0.28, size=12.3, color=RED, bold=True, align=PP_ALIGN.CENTER)
finish(slide, note)


# 43
note = (
    "The final 48-hour process is a controlled deployment. Verify checksums and labels, generate three pre-frozen candidates, run structural and distributional QC, "
    "apply the predeclared selection rule, and submit with recovery margin. Do not infer hidden cell-line identities."
)
slide = new_slide("Competition operations", "Final-phase 48-hour playbook",
                  "Treat D/E/F ingestion as production adaptation—not a new research sprint.",
                  "Team final-phase recommendation", note)
steps = [
    ("0–2 h", "Ingest", "checksum • manifest • D/E/F labels", CYAN),
    ("2–6 h", "Profile", "depth • states • similarity • uncertainty", PURPLE),
    ("6–18 h", "Adapt", "context encoder + frozen calibration rule", PINK),
    ("18–30 h", "Generate", "baseline • generator • ensemble", GREEN),
    ("30–38 h", "QC", "counts • axes • NNZ • moments • DE calls", YELLOW),
    ("38–44 h", "Select", "predeclared robust decision rule", BLUE),
    ("44–48 h", "Package", "dry run • hash • upload • archive", RED),
]
for i, (time, ttl, body, col) in enumerate(steps):
    y = 1.60 + i * 0.70
    circle(slide, 0.78, y + 0.04, 0.46, fill=col)
    text(slide, str(i + 1), 0.78, y + 0.04, 0.46, 0.46, size=10, color=BG if col in [CYAN, GREEN, YELLOW] else WHITE,
         bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    text(slide, time, 1.46, y + 0.05, 1.02, 0.26, size=10.2, color=col, bold=True, font=MONO)
    text(slide, ttl, 2.62, y + 0.03, 1.26, 0.30, size=12.2, bold=True)
    text(slide, body, 4.08, y + 0.03, 4.38, 0.30, size=10.5, color=MUTED)
    if i < 6:
        line(slide, 1.01, y + 0.50, 1.01, y + 0.74, color=MUTED2, width=1.5)
rect(slide, 8.82, 1.78, 3.52, 3.46, fill=SURFACE, line=RED)
text(slide, "Forbidden changes", 9.14, 2.08, 2.88, 0.34, size=15, color=RED, bold=True, align=PP_ALIGN.CENTER)
bullet_list(slide, [
    "new architecture",
    "new loss family",
    "manual cell-line identity guess",
    "unreviewed target mapping",
    "ad hoc per-target editing",
    "deadline-day first rehearsal"
], 9.12, 2.64, 2.92, 1.92, size=10.8, bullet_color=RED, gap=3)
rect(slide, 8.82, 5.52, 3.52, 0.76, fill=BG2, line=GREEN)
text(slide, "Preferred upload: 4 Nov", 9.08, 5.73, 3.00, 0.28, size=12.2, color=GREEN2, bold=True, align=PP_ALIGN.CENTER)
finish(slide, note)


# 44
note = (
    "End the discussion with decisions, not broad agreement. Assign names to P1/P2/P3, define the first baseline and data sources, choose a metric-emulator owner, "
    "set compute budget, and freeze the next review date."
)
slide = new_slide("Discussion", "Decisions to make today",
                  "Leave the meeting with owners, gates, and a two-week definition of done.",
                  None, note)
decisions = [
    ("01", "Who is P1, P2, and P3?", "Name one accountable owner and one backup per critical lane.", CYAN),
    ("02", "Which external datasets are legal and useful?", "Create the provenance registry before downloading everything.", PURPLE),
    ("03", "What is the strongest baseline?", "Agree on control, global, target-transfer, and context-interaction rungs.", PINK),
    ("04", "Who owns exact metric parity?", "This owner also controls LOCO folds and score reports.", GREEN),
    ("05", "What is the diffusion go/no-go gate?", "Require held-out benefit beyond NB/DM count sampling.", YELLOW),
    ("06", "What is the RL entry gate?", "Cap compute and demand a strong supervised generator first.", BLUE),
]
for i, (num, ttl, body, col) in enumerate(decisions):
    x = 0.68 + (i % 2) * 6.12
    y = 1.68 + (i // 2) * 1.54
    rect(slide, x, y, 5.78, 1.22, fill=SURFACE, line=col)
    circle(slide, x + 0.20, y + 0.25, 0.66, fill=col)
    text(slide, num, x + 0.20, y + 0.25, 0.66, 0.66, size=11, color=BG if col in [CYAN, GREEN, YELLOW] else WHITE,
         bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    text(slide, ttl, x + 1.08, y + 0.20, 4.34, 0.32, size=12.4, bold=True)
    text(slide, body, x + 1.08, y + 0.62, 4.34, 0.36, size=9.8, color=MUTED)
text(slide, "Next review: inspect the current 18 metric cells and approve the first trusted LOCO baseline.",
     1.22, 6.46, 10.88, 0.32, size=13.2, color=CYAN2, bold=True, align=PP_ALIGN.CENTER)
finish(slide, note)


# 45
note = (
    "This is the recommended minimum viable scientific system. It is intentionally simpler than a full diffusion model. It provides a trustworthy baseline, "
    "a biologically meaningful effect predictor, calibrated raw-count cells, and a complete audit trail."
)
slide = new_slide("Recommendation", "The first build we should ship",
                  "A minimum viable scientific system—not merely a minimum viable file.",
                  "Team recommendation", note)
rect(slide, 0.72, 1.70, 7.56, 4.42, fill=SURFACE, line=PURPLE)
text(slide, "MVP architecture", 1.06, 2.00, 2.38, 0.34, size=16, color=PURPLE2, bold=True)
bullet_list(slide, [
    "Context: robust NTC pseudobulk + gene-module activities + depth statistics.",
    "Target: frozen multi-view gene embedding plus trainable low-rank adapter.",
    "Effect: low-rank Δ predictor with global and context-interaction terms.",
    "Losses: centroid, direction/rank, LFC calibration, and shrinkage.",
    "Cells: real-control anchoring + empirical depth + NB/DM dispersion.",
    "Selection: exact six-metric LOCO score with mean and worst-context gates."
], 1.08, 2.56, 6.66, 2.94, size=12.0, bullet_color=CYAN, gap=4)
rect(slide, 8.66, 1.70, 3.72, 4.42, fill=BG2, line=GREEN)
text(slide, "Acceptance gates", 8.98, 2.00, 3.08, 0.34, size=15, color=GREEN2, bold=True, align=PP_ALIGN.CENTER)
gates = [("G0", "valid 360k package"), ("G1", "beats statistical baseline"), ("G2", "worst context protected"),
         ("G3", "three-seed stability"), ("G4", "one-command regeneration")]
for i, (gid, body) in enumerate(gates):
    y = 2.60 + i * 0.60
    label(slide, gid, 9.04, y, 0.62, fill=GREEN, color=BG, size=8.5)
    text(slide, body, 9.84, y + 0.02, 2.10, 0.24, size=10.2, bold=True)
text(slide, "Only after G0–G4: test residual diffusion; only after that gate: consider RL.",
     1.20, 6.43, 10.92, 0.30, size=13, color=YELLOW2, bold=True, align=PP_ALIGN.CENTER)
finish(slide, note)


# 46
note = (
    "Use the primary sources for rules and biology. The deck distinguishes official facts, measured local evidence, and team recommendations. "
    "Access dates are 30 August 2026 for web documentation."
)
slide = new_slide("References", "Primary sources and local evidence",
                  "Official facts, scientific background, and reproducible local measurements.",
                  "Accessed 30 August 2026", note)
refs_left = [
    "VCC 2026 Data\nvirtualcellchallenge.org/datasets",
    "VCC 2026 Evaluation\nvirtualcellchallenge.org/evaluation",
    "Official Rules\nvirtualcellchallenge.org/rules",
    "VCC CLI Guide\nvcc-cli-wiki.virtualcellchallenge.org",
    "Metric specification\ngithub.com/ArcInstitute/cell-eval2/…/vcc2026-metrics.pdf",
]
refs_right = [
    "Qi et al. CRISPRi. Cell (2013)\ndoi.org/10.1016/j.cell.2013.02.022",
    "Dixit et al. Perturb-seq. Cell (2016)\ndoi.org/10.1016/j.cell.2016.11.038",
    "Lotfollahi et al. CPA. Mol Syst Biol (2023)\ndoi.org/10.15252/msb.202211517",
    "Roohani et al. GEARS. Nat Biotechnol (2024)\ndoi.org/10.1038/s41587-023-01905-6",
    "Local evidence\nartifacts/*.json • reports/* • Slurm job 792580",
]
for i, ref in enumerate(refs_left):
    y = 1.70 + i * 0.94
    circle(slide, 0.80, y + 0.10, 0.42, fill=CYAN)
    text(slide, str(i + 1), 0.80, y + 0.10, 0.42, 0.42, size=9, color=BG, bold=True,
         align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    text(slide, ref, 1.44, y, 4.80, 0.68, size=10.5, color=WHITE, bold=True)
for i, ref in enumerate(refs_right):
    y = 1.70 + i * 0.94
    circle(slide, 6.90, y + 0.10, 0.42, fill=PURPLE)
    text(slide, str(i + 6), 6.90, y + 0.10, 0.42, 0.42, size=9, color=WHITE, bold=True,
         align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    text(slide, ref, 7.54, y, 4.84, 0.68, size=10.5, color=WHITE, bold=True)
rect(slide, 1.06, 6.50, 11.18, 0.28, fill=BG2, line=SURFACE2)
text(slide, "Evidence labels used throughout: official fact • measured local result • team recommendation",
     1.32, 6.54, 10.66, 0.18, size=9.6, color=MUTED, bold=True, align=PP_ALIGN.CENTER)
finish(slide, note)


def write_notes():
    lines = [
        "# VCC 2026 — English Speaker Notes",
        "",
        "Companion notes for `VCC_2026_Three_Person_Team_Strategy_EN.pptx`.",
        "All factual competition statements use the 30 August 2026 evidence snapshot.",
        "",
    ]
    for n, title_value, note_value in NOTES:
        lines.extend([f"## Slide {n:02d} — {title_value}", "", note_value.strip(), ""])
    lines.extend([
        "## Suggested presenter split", "",
        "- P1 (Biological Modeling): slides 1–15, 18–27, 44–46.",
        "- P2 (Generative & Optimization): slides 28–34 and model questions.",
        "- P3 (Evaluation, MLOps & Competition): slides 16–17, 35–43.",
        "- For a 20-minute talk, use slides 1, 2, 4, 5, 8, 11–13, 16–19, 23, 28, 33–35, 37, 39–41, 45.",
        "",
    ])
    NOTES_OUT.write_text("\n".join(lines), encoding="utf-8")


write_notes()
OUT.parent.mkdir(parents=True, exist_ok=True)
prs.save(OUT)
print(f"saved={OUT}")
print(f"slides={len(prs.slides)}")
print(f"notes={NOTES_OUT}")
